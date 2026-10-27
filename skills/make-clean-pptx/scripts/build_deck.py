#!/usr/bin/env python3
"""Build a .pptx by cloning slides from the design / shapes templates and filling them.

    build_deck.py --spec deck.json [--out deck.pptx] [--local-dir DIR]
                  [--design PPTX --shapes PPTX] [--strict]

Why cloning: the templates already *are* the design. Re-drawing slides from
scratch can only approximate them, cloning reproduces them exactly (theme,
fonts, spacing, decorative shapes, charts) and we only swap the content.

Deck spec (JSON)
----------------
{
  "output": "out/deck.pptx",
  "properties": {"title": "Q3 review", "author": "Jane Doe"},
  "slides": [
    {"source": "design", "slide": 1,
     "fill": {"Title": "A **bold** word", "Subtitle": "Q3 2026"},
     "notes": "Speaker notes"},

    {"source": "design", "slide": 14,
     "fill": {"Title": "Revenue by region",
              "Body": ["First point", {"text": "Sub point", "level": 1}],
              "Photo": {"path": "img/team.jpg", "alt": "Team at work", "credit": "..."},
              "Extra shape": null}},

    {"source": "design", "slide": 22,
     "fill": {"Title": "Roadmap"},
     "insert": [{"source": "shapes", "slide": 5, "component": "Timeline 4 steps",
                 "into": "Content Area",
                 "fill": {"Step 1 title": "Discover", "Step 2 title": "Build"}}]}
  ]
}

fill values
  str / list     text (paragraphs; "**x**" = bold; list items may be {"text","level"})
  {"path": ...}  picture (cover-crop by default; "fit": "contain" keeps all of it)
  {"categories": [...], "series": {"Name": [...]}}   chart data
  [[...], [...]] table cells (rows are added or removed to fit)
  null or ""     delete the shape (use for optional slots you do not need)

"logo" (optional) puts a PNG, JPG or SVG file in the "Logo" slot of the covers
and the closing slide, scaled to fit with its aspect ratio kept (path relative to
the spec file, then to the current folder); "logo_alt" is its alt text, usually
the brand name (without it the logo is marked decorative). Without "logo" the
slots are removed, as any optional field left unfilled.
Keys are shape names; "#3" means the third top-level shape; if no shape has that
name, "{{key}}" tokens inside text are replaced instead.

Prints a JSON report; warnings list anything the agent should fix (unknown shape
names, text that probably overflows, unreadable images...).
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from copy import deepcopy
from pathlib import Path

try:
    from lxml import etree
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.oxml.ns import qn

    from _clone import SlideCloner, delete_slide
    from _common import eprint, estimate_fit, is_dark, iter_shapes, mark_decorative, overflows, shape_kind, token_xml
    from _svg import SvgError, parse_svg, svg_to_shapes
    from fetch_templates import FetchError, ensure_templates
    from read_index import load_index, resolve
except ImportError as _exc:  # a missing package would otherwise end in a traceback
    if __name__ != "__main__":
        raise
    print(json.dumps({"error": f"a Python package the skill needs is missing ({_exc.name})",
                      "hint": "Run scripts/doctor.py --fix to install it, then build again."}))
    sys.exit(1)

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"}
SOURCE_ALIASES = {"design": "layouts", "shapes": "visuals"}
SPEC_KEYS = {"output", "properties", "slides", "language", "logo", "logo_alt"}
LOGO_FIELD = "Logo"  # the logo slot of the layouts (role=logo in the index)
LOGO_EXT = {".png", ".jpg", ".jpeg", ".svg"}
SLIDE_KEYS = {"source", "slide", "id", "fill", "insert", "notes", "keep_unfilled", "hidden"}
INSERT_KEYS = {"source", "slide", "id", "component", "into", "box", "fit", "fill", "keep_unfilled"}


def check_keys(spec: dict) -> list:
    """Typos in the spec are silent otherwise ("fil" instead of "fill"): report them."""
    import difflib

    out = []

    def unknown(obj, allowed, where):
        for k in obj:
            if k not in allowed:
                near = difflib.get_close_matches(k, sorted(allowed), n=1)
                out.append(f"{where}: unknown key '{k}' ignored" + (f" (did you mean '{near[0]}'?)" if near else ""))

    unknown(spec, SPEC_KEYS, "spec")
    for i, s in enumerate(spec.get("slides") or []):
        if isinstance(s, dict):
            unknown(s, SLIDE_KEYS, f"slides[{i}]")
            for k, c in enumerate(s.get("insert") or []):
                if isinstance(c, dict):
                    unknown(c, INSERT_KEYS, f"slides[{i}].insert[{k}]")
    return out


MIN_SCALED_PT = 8


class SpecError(ValueError):
    pass


def spec_path(value, spec_dir: Path) -> Path:
    """A file named in the spec: a relative path is looked up next to the spec file, then in the current folder."""
    p = Path(str(value)).expanduser()
    if p.is_absolute():
        return p
    for base in (spec_dir, Path.cwd()):
        if (base / p).exists():
            return base / p
    return spec_dir / p


def load_logo(spec: dict, spec_dir: Path) -> dict | None:
    """Check the deck's "logo" before any work: a wrong file stops the build with a plain message."""
    value = spec.get("logo")
    if value in (None, ""):
        return None
    if not isinstance(value, str):
        raise SpecError("'logo' must be the path of a PNG, JPG or SVG file")
    path = spec_path(value, spec_dir)
    if path.suffix.lower() not in LOGO_EXT:
        raise SpecError(f"logo '{value}': use a PNG, JPG or SVG file (got '{path.suffix or 'no extension'}')")
    if not path.is_file():
        raise SpecError(f"logo '{value}' not found (looked in {spec_dir} and in the current folder)")
    if path.suffix.lower() == ".svg":
        try:
            parse_svg(path.read_text(encoding="utf-8"))
        except (SvgError, UnicodeDecodeError) as exc:
            raise SpecError(f"logo '{value}' is not a usable SVG ({exc}); export it as PNG instead") from exc
    else:
        from PIL import Image

        try:
            with Image.open(path) as im:
                fmt = im.format
                im.verify()
        except Exception as exc:  # noqa: BLE001 - Pillow raises many kinds of errors on bad files
            raise SpecError(f"logo '{value}' is not a readable image ({exc.__class__.__name__})") from exc
        if fmt not in ("PNG", "JPEG"):
            raise SpecError(f"logo '{value}' is a {fmt} file; use a PNG, JPG or SVG file")
    alt = spec.get("logo_alt") or ""
    if not isinstance(alt, str):
        raise SpecError("'logo_alt' must be text (the brand name)")
    return {"path": str(path), "alt": alt.strip()}


class Builder:
    def __init__(self, design_path: str, shapes_path: str | None, design_index=None, shapes_index=None, strict=False,
                 logo: dict | None = None, spec_dir: Path | None = None):
        self.prs = Presentation(design_path)
        self.logo = logo
        self.logos_placed = 0
        self.spec_dir = spec_dir or Path.cwd()
        self.shapes_prs = Presentation(shapes_path) if shapes_path else None
        self.design_index = design_index or []
        self.shapes_index = shapes_index or []
        self.strict = strict
        self.cloner = SlideCloner(self.prs)
        self.warnings: list[str] = []
        self.credits: list[dict] = []
        self._n_original = len(self.prs.slides)
        if self.shapes_prs and (
            self.shapes_prs.slide_width != self.prs.slide_width or self.shapes_prs.slide_height != self.prs.slide_height
        ):
            self.warn("Design and shapes templates have different slide sizes; inserted components are rescaled, "
                      "whole shapes-template slides will not match the design.")

    # ----------------------------------------------------------------- utils
    def warn(self, msg: str) -> None:
        self.warnings.append(msg)

    def _source_deck(self, name: str):
        name = SOURCE_ALIASES.get(name, name)
        if name == "layouts":
            return self.prs, self.design_index
        if name == "visuals":
            if not self.shapes_prs:
                raise SpecError("spec uses source 'visuals' but no visuals library was loaded")
            return self.shapes_prs, self.shapes_index
        raise SpecError(f"unknown source '{name}' (use 'layouts' or 'visuals')")

    def _resolve_slide_no(self, item: dict, index: list[dict], where: str) -> int:
        if "slide" in item:
            return int(item["slide"])
        if "id" in item:
            hit = resolve(index, id_=item["id"])
            if hit:
                return hit["slide"]
            raise SpecError(f"{where}: id '{item['id']}' not found in the index")
        raise SpecError(f"{where}: needs 'slide' (number) or 'id'")

    # ----------------------------------------------------------------- build
    def add_slide(self, spec: dict, pos: int) -> None:
        where = f"slides[{pos}]"
        src_name = SOURCE_ALIASES.get(spec.get("source", "layouts"), spec.get("source", "layouts"))
        deck, index = self._source_deck(src_name)
        n = self._resolve_slide_no(spec, index, where)
        if src_name == "layouts":
            limit = self._n_original
        else:
            limit = len(deck.slides)
        if not 1 <= n <= limit:
            raise SpecError(f"{where}: {src_name} slide {n} does not exist (1..{limit})")
        src_slide = deck.slides[n - 1]
        slide = self.cloner.clone_slide(src_slide)
        label = f"slide {pos + 1} ({src_name}#{n})"

        fill = dict(spec.get("fill") or {})
        if (self.logo and not any(str(k).strip().lower() == LOGO_FIELD.lower() for k in fill)
                and any(s.name == LOGO_FIELD for s in slide.shapes)):
            fill[LOGO_FIELD] = dict(self.logo)  # the deck's logo goes in every logo slot
        self._fill(slide, list(iter_shapes(slide)), list(slide.shapes), fill, label)

        used = set(fill)
        for k, comp in enumerate(spec.get("insert") or []):
            self._insert_component(slide, comp, f"{label} insert[{k}]")
            if comp.get("into"):
                used.add(comp["into"])
        if not spec.get("keep_unfilled"):
            self._clear_unfilled(list(slide.shapes), resolve(index, slide=n), used, label)

        notes = spec.get("notes")
        credit_lines = [c["credit"] for c in self.credits if c.get("slide") == pos + 1 and c.get("credit")]
        if credit_lines:
            notes = ((notes + "\n\n") if notes else "") + "Image credits: " + " | ".join(credit_lines)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    def finish(self, props: dict | None, language: str | None = None) -> None:
        for _ in range(self._n_original):
            delete_slide(self.prs, 0)
        if language:
            set_language(self.prs, language)
        cp = self.prs.core_properties
        props = props or {}
        cp.title = props.get("title", "")
        cp.author = props.get("author", "")
        cp.subject = props.get("subject", "")
        cp.keywords = props.get("keywords", "")
        cp.comments = props.get("comments", "")
        cp.last_modified_by = props.get("author", "")
        self.warnings.extend(self.cloner.warnings)
        if self.logo and not self.logos_placed:
            self.warn("spec: 'logo' is not shown: no slide has a logo slot (covers L-01 to L-05, closing L-49)")

    # ---------------------------------------------------------------- insert
    def _insert_component(self, slide, comp: dict, label: str) -> None:
        deck, index = self._source_deck(comp.get("source", "shapes"))
        n = self._resolve_slide_no(comp, index, label)
        if not 1 <= n <= len(deck.slides):
            raise SpecError(f"{label}: slide {n} does not exist")
        src_slide = deck.slides[n - 1]
        name = comp.get("component")
        if not name:
            raise SpecError(f"{label}: 'component' (top-level shape or group name) is required")
        matches = [s for s in src_slide.shapes if s.name == name] or [
            s for s in src_slide.shapes if s.name.strip().lower() == name.strip().lower()
        ]
        if not matches:
            names = ", ".join(repr(s.name) for s in src_slide.shapes)
            raise SpecError(f"{label}: no top-level shape named '{name}' on that slide. Available: {names}")

        # target box in EMU
        sw, sh = self.prs.slide_width, self.prs.slide_height
        into = comp.get("into")
        if into:
            frames = [s for s in iter_shapes(slide) if s.name == into]
            if not frames:
                raise SpecError(f"{label}: 'into' shape '{into}' not found on the destination slide")
            f = frames[0]
            box = (f.left, f.top, f.width, f.height)
            f._element.getparent().remove(f._element)
        elif comp.get("box"):
            x, y, w, h = comp["box"]
            box = (int(sw * x / 100), int(sh * y / 100), int(sw * w / 100), int(sh * h / 100))
        else:
            box = None

        shp = self.cloner.clone_shape(src_slide, matches[0], slide)
        if box and shp.width and shp.height:
            bx, by, bw, bh = box
            scale = min(bw / shp.width, bh / shp.height)
            if comp.get("fit") == "stretch":
                new_w, new_h = bw, bh
                scale = min(bw / shp.width, bh / shp.height)
            else:
                new_w, new_h = int(shp.width * scale), int(shp.height * scale)
            shp.left = int(bx + (bw - new_w) / 2)
            shp.top = int(by + (bh - new_h) / 2)
            shp.width, shp.height = new_w, new_h
            if scale < 0.97:
                clamped = _scale_fonts(shp._element, scale)
                if clamped:
                    self.warn(f"{label}: '{name}' is shown at {scale:.0%} of its size; {clamped} text runs hit the "
                              f"{MIN_SCALED_PT}pt floor. Prefer a larger frame (see the 'frame' column) or shorter labels.")

        scope = [shp] + (list(iter_shapes(shp)) if shape_kind(shp) == "group" else [])
        cfill = comp.get("fill") or {}
        self._fill(slide, scope, None, cfill, label)
        if not comp.get("keep_unfilled"):
            entry = resolve(index, id_=comp["id"]) if comp.get("id") else None
            entry = entry or next((e for e in index if e["slide"] == n and e["name"] == name), None) \
                or next((e for e in index if e["slide"] == n), None)
            self._clear_unfilled(scope, entry, set(cfill), label)

    # ------------------------------------------------------------------ fill
    def _fill(self, slide, scope: list, top_level: list | None, fill: dict, label: str) -> None:
        for key, value in fill.items():
            targets = self._find(scope, top_level, key)
            if not targets:
                if isinstance(value, (str, list)) and self._replace_token(scope, key, value):
                    continue
                self._unknown(key, scope, label)
                continue
            if len(targets) > 1:
                self.warn(f"{label}: key '{key}' matches {len(targets)} shapes; filling all of them")
            for shp in targets:
                try:
                    self._apply(slide, shp, key, value, label)
                except Exception as exc:  # noqa: BLE001
                    self.warn(f"{label}: could not fill '{key}': {exc}")
                    if self.strict:
                        raise

    def _unknown(self, key, scope, label):
        names = sorted({s.name for s in scope})
        msg = f"{label}: no shape named '{key}'. Shapes here: {', '.join(names)}"
        self.warn(msg)
        if self.strict:
            raise SpecError(msg)

    @staticmethod
    def _find(scope, top_level, key: str) -> list:
        if key.startswith("#") and key[1:].isdigit() and top_level is not None:
            i = int(key[1:]) - 1
            return [top_level[i]] if 0 <= i < len(top_level) else []
        exact = [s for s in scope if s.name == key]
        if exact:
            return exact
        low = key.strip().lower()
        return [s for s in scope if s.name.strip().lower() == low]

    def _apply(self, slide, shp, key, value, label):
        kind = shape_kind(shp)
        if value is None or value == "":
            shp._element.getparent().remove(shp._element)
            return
        if shp.name == LOGO_FIELD and isinstance(value, str) and Path(value).suffix.lower() in LOGO_EXT:
            value = {"path": value}
        if shp.name == LOGO_FIELD and isinstance(value, dict) and "path" in value:
            self._fill_logo(slide, shp, value, label)
            return
        if isinstance(value, dict) and ("icon" in value or "svg" in value):
            self._fill_svg(slide, shp, value, label)
        elif isinstance(value, dict) and "path" in value and str(value["path"]).lower().endswith(".svg"):
            self._fill_svg(slide, shp, dict(value, svg=value["path"], mode=value.get("mode", "grey")), label)
        elif isinstance(value, dict) and "path" in value:
            self._fill_image(slide, shp, value, label)
        elif isinstance(value, dict) and ("series" in value or "waterfall" in value or "gantt" in value):
            if kind != "chart":
                raise ValueError("chart data given but the shape is not a chart")
            self._fill_chart(shp, value)
        elif isinstance(value, dict) and ({"text", "fill", "color", "at"} & set(value)):
            self._fill_rich(slide, shp, value, label)
        elif isinstance(value, list) and value and isinstance(value[0], list):
            if kind != "table":
                raise ValueError("table rows given but the shape is not a table")
            self._fill_table(shp, value)
        elif isinstance(value, (str, list, int, float)):
            if isinstance(value, (int, float)):
                value = str(value)
            if kind == "picture" and isinstance(value, str) and Path(value).suffix.lower() in IMAGE_EXT:
                self._fill_image(slide, shp, {"path": value}, label)
                return
            if not getattr(shp, "has_text_frame", False):
                raise ValueError(f"shape is a {kind}, it cannot hold text")
            paras = set_text_frame(shp.text_frame, value)
            fit = estimate_fit(shp, paras)
            if overflows(fit):
                self.warn(
                    f"{label}: '{shp.name}' will probably overflow (needs ~{fit['needed_lines']} lines, "
                    f"room for ~{fit['avail_lines']} at {fit['font_pt']}pt). Shorten the text or pick a roomier slide."
                )
        else:
            raise ValueError(f"unsupported value type {type(value).__name__}")

    # ------------------------------------------------------- unfilled fields
    def _clear_unfilled(self, scope: list, entry: dict | None, used: set, label: str) -> None:
        """Delete optional slots that were not filled; warn about required ones.

        Template text is a description of the slot, never content, so it must not
        reach the final deck.
        """
        if not entry:
            return
        used_low = {str(k).strip().lower() for k in used}

        def alive(name):
            return [s for s in scope if s.name == name and s._element.getparent() is not None]

        fields = entry.get("field_list", [])
        # 1. plain fields: delete unfilled optional slots, warn about required ones
        for f in fields:
            if f["type"] in ("plot", "deco") or f.get("follows") or f["name"].strip().lower() in used_low:
                continue
            targets = alive(f["name"])
            if not targets or f.get("default"):
                continue
            if f["optional"]:
                for t in targets:
                    t._element.getparent().remove(t._element)
            else:
                what = {"frame": "insert a visual into it", "chart": "give it data", "image": "give it a picture",
                        "table": "give it rows", "icon": "give it an icon"}.get(f["type"], "fill it")
                self.warn(f"{label}: required field '{f['name']}' still shows its template placeholder; {what} or set it to null.")
        # 2. companions (numbers, dots, rules) go with the field they belong to
        for f in fields:
            partner = f.get("follows")
            if not partner or f["name"].strip().lower() in used_low:
                continue
            if not alive(partner):
                for t in alive(f["name"]):
                    t._element.getparent().remove(t._element)

    # ------------------------------------------------------------ rich text
    def _fill_rich(self, slide, shp, value: dict, label: str) -> None:
        """{"text": "...", "fill": "a1", "color": "paper", "at": [x%, y%]}"""
        if "text" in value:
            self._apply(slide, shp, shp.name, value["text"] if value["text"] is not None else "", label)
        if "fill" in value:
            _set_fill(shp._element, value["fill"])
            if "color" not in value and getattr(shp, "has_text_frame", False):
                _set_text_colour(shp._element, "paper" if is_dark(value["fill"]) else "ink")
        if "color" in value:
            _set_text_colour(shp._element, value["color"])
        if "at" in value:
            parent = shp._element.getparent()
            plots = [s for s in parent.iterchildren() if s.xpath("./*/p:cNvPr[@name='Plot Area']")]
            if not plots:
                raise ValueError("'at' needs a shape named 'Plot Area' next to this one (matrix and map visuals)")
            from pptx.shapes.autoshape import Shape

            plot = Shape(plots[0], None)
            fx, fy = (max(0.0, min(100.0, float(v))) / 100 for v in value["at"])
            shp.left = int(plot.left + fx * plot.width - shp.width / 2)
            shp.top = int(plot.top + (1 - fy) * plot.height - shp.height / 2)

    # ---------------------------------------------------------- icons / svg
    def _fill_svg(self, slide, shp, value: dict, label: str) -> None:
        """Replace a slot with an SVG converted to native shapes.

        {"icon": "lucide:rocket"} or {"svg": "path/file.svg"}; options "mode"
        (mono, grey, keep), "color" (theme token) and "weight" (stroke factor).
        """
        meta = {}
        if "icon" in value:
            from find_icon import IconError, get_svg

            try:
                text, meta = get_svg(str(value["icon"]))
            except IconError as exc:
                raise ValueError(str(exc)) from exc
            mode = value.get("mode", "mono")
        else:
            path = Path(str(value["svg"])).expanduser()
            if path.suffix.lower() != ".svg":
                raise ValueError(f"'svg' expects an .svg file, got {path.name}")
            if not path.is_file():
                raise FileNotFoundError(f"svg not found: {path}")
            text = path.read_text(encoding="utf-8")
            mode = value.get("mode", "grey")
        if value.get("color"):
            mono = token_xml(value["color"])
        else:
            found = shp._element.xpath("./p:spPr/a:ln/a:solidFill/*")
            mono = etree.tostring(found[0]).decode() if found else token_xml("ink")
        size_in = min(shp.width, shp.height) / 914400
        weight = float(value.get("weight", 1.0 if size_in <= 0.7 else 0.8))
        ids = [int(i) for i in slide._element.xpath("//p:cNvPr/@id") if str(i).isdigit()]
        try:
            xml, info = svg_to_shapes(text, shp.left, shp.top, shp.width, shp.height, name=shp.name, mode=mode,
                                      mono_xml=mono, start_id=max(ids or [1]) + 1, descr=value.get("alt", ""), weight=weight)
        except SvgError as exc:
            raise ValueError(f"cannot convert this SVG: {exc}") from exc
        if info["skipped"]:
            self.warn(f"{label}: '{shp.name}': SVG parts not supported and skipped ({', '.join(info['skipped'])}); check the render.")
        new = etree.fromstring(xml)
        shp._element.addprevious(new)
        shp._element.getparent().remove(shp._element)
        if meta.get("attribution_required"):
            self.credits.append({"slide": int(re.search(r"slide (\d+)", label).group(1)), "credit": meta["credit"], "path": meta.get("ref", "")})

    # ---------------------------------------------------------------- logo
    def _fill_logo(self, slide, shp, value: dict, label: str) -> None:
        """The logo, whole, in its slot: scaled to fit, aspect ratio kept, aligned on the slot's side.

        Alt text is the brand name; without one the logo is marked decorative (it repeats
        what the slide already says). The result keeps the name "Logo" so tools can find it.
        """
        path = spec_path(value["path"], self.spec_dir)
        if path.suffix.lower() not in LOGO_EXT:
            raise ValueError(f"logo: use a PNG, JPG or SVG file, not {path.name}")
        if not path.is_file():
            raise FileNotFoundError(f"logo not found: {path}")
        alt = str(value.get("alt") or "").strip()
        slot = (shp.left, shp.top, shp.width, shp.height)
        parent = shp._element.getparent()
        at = parent.index(shp._element)
        if path.suffix.lower() == ".svg":
            self._fill_svg(slide, shp, {"svg": str(path), "mode": value.get("mode", "keep"), "alt": alt}, label)
        else:
            self._fill_image(slide, shp, {"path": str(path), "fit": "contain", "alt": alt, "decorative": not alt}, label)
        new = parent[at]
        cnvpr = new[0].find(qn("p:cNvPr"))  # nvSpPr, nvPicPr or nvGrpSpPr
        cnvpr.set("name", LOGO_FIELD)
        if not alt:
            mark_decorative(cnvpr)
        _align_in_slot(new, slot, self.prs.slide_width)
        self.logos_placed += 1

    @staticmethod
    def _replace_token(scope, key, value) -> bool:
        token = "{{" + key + "}}"
        text = "\n".join(value) if isinstance(value, list) else str(value)
        done = False
        for shp in scope:
            if not getattr(shp, "has_text_frame", False):
                continue
            for p in shp.text_frame.paragraphs:
                joined = "".join(r.text for r in p.runs)
                if token in joined and p.runs:
                    p.runs[0].text = joined.replace(token, text)
                    for r in p.runs[1:]:
                        r.text = ""
                    done = True
        return done

    # --------------------------------------------------------------- images
    def _fill_image(self, slide, shp, spec: dict, label: str) -> None:
        from PIL import Image

        path = Path(spec["path"]).expanduser()
        if not path.is_file():
            raise FileNotFoundError(f"image not found: {path}")
        with Image.open(path) as im:
            iw, ih = im.size
        fit = spec.get("fit", "cover")
        alt = spec.get("alt", "")

        left, top, width, height = shp.left, shp.top, shp.width, shp.height
        el = shp._element
        is_pic_element = el.tag == qn("p:pic")

        if hasattr(shp, "insert_picture") and not is_pic_element:
            pic = shp.insert_picture(str(path))  # crops to fill the placeholder
        elif is_pic_element:
            image_part, rId = slide.part.get_or_add_image_part(str(path))
            old = el.blipFill.blip.rEmbed
            el.blipFill.blip.rEmbed = rId
            if old and old != rId:
                slide.part.drop_rel(old)
            pic = shp
            _apply_crop(pic, iw, ih, width, height, fit)
            if fit == "contain":
                _contain(pic, left, top, width, height, iw, ih)
        else:
            pic = slide.shapes.add_picture(str(path), left, top, width, height)
            _apply_crop(pic, iw, ih, width, height, fit)
            if fit == "contain":
                _contain(pic, left, top, width, height, iw, ih)
            parent = el.getparent()
            parent.insert(parent.index(el), pic._element)  # keep z-order
            parent.remove(el)

        pic._element.xpath(".//p:cNvPr")[0].set("descr", alt)
        if not alt and not spec.get("decorative"):
            self.warn(f"{label}: image for '{shp.name}' has no 'alt' text (accessibility).")
        if spec.get("credit"):
            self.credits.append({"slide": int(re.search(r"slide (\d+)", label).group(1)), "credit": spec["credit"],
                                 "path": str(path)})

    # ------------------------------------------------------------- chart/table
    @staticmethod
    def _fill_chart(shp, value: dict) -> None:
        if "waterfall" in value:
            value = dict(value, series=waterfall_series(value["waterfall"], value.get("totals")))
        elif "gantt" in value:
            starts, durations = zip(*[(float(a), float(b)) for a, b in value["gantt"]]) if value["gantt"] else ((), ())
            value = dict(value, series={"Start": list(starts), "Duration": list(durations)})
        cd = CategoryChartData()
        cd.categories = value.get("categories", [])
        for name, vals in value["series"].items():
            cd.add_series(name, vals)
        shp.chart.replace_data(cd)

    def _fill_table(self, shp, rows: list[list]) -> None:
        tbl = shp.table
        tbl_el = tbl._tbl
        trs = tbl_el.findall(qn("a:tr"))
        while len(trs) < len(rows):
            new_tr = deepcopy(trs[-1])
            trs[-1].addnext(new_tr)
            trs = tbl_el.findall(qn("a:tr"))
        for extra in trs[len(rows):]:
            tbl_el.remove(extra)
        for r, row in enumerate(rows):
            cells = tbl.rows[r].cells
            for c, val in enumerate(row):
                if c < len(cells):
                    set_text_frame(cells[c].text_frame, "" if val is None else str(val))
            if len(row) > len(cells):
                self.warn(f"table '{shp.name}': row {r + 1} has {len(row)} values but the table has {len(cells)} columns")


# --------------------------------------------------------------------------
# Charts, colours, scaling
# --------------------------------------------------------------------------
def set_language(prs, language: str) -> None:
    """Tag every text run with the deck language (screen readers, spell check)."""
    if not re.match(r"^[A-Za-z]{2,3}(-[A-Za-z0-9]{2,8})*$", language):
        raise SpecError(f"language '{language}' is not a language tag such as fr-FR or en-GB")
    tags = (qn("a:rPr"), qn("a:endParaRPr"), qn("a:defRPr"))
    for slide in prs.slides:
        parts = [slide._element] + ([slide.notes_slide._element] if slide.has_notes_slide else [])
        for el in parts:
            for node in el.iter(*tags):
                node.set("lang", language)
    prs.core_properties.language = language


def waterfall_series(values: list, totals: list | None = None) -> dict:
    """[start, +a, -b, ..., end] -> the four series of the waterfall visual.

    Indexes in `totals` (default: first and last) are absolute bars; the others
    are steps from the running total. Series order matches the template:
    Base (hidden), Total, Increase, Decrease.
    """
    n = len(values)
    totals = set(range(n) if n < 2 else (totals if totals is not None else [0, n - 1]))
    base, tot, inc, dec = [0.0] * n, [0.0] * n, [0.0] * n, [0.0] * n
    running = 0.0
    for i, v in enumerate(float(x) for x in values):
        if i in totals:
            tot[i] = v
            running = v
        elif v >= 0:
            base[i], inc[i] = running, v
            running += v
        else:
            running += v
            base[i], dec[i] = running, -v
    return {"Base": base, "Total": tot, "Increase": inc, "Decrease": dec}


def _set_fill(el, colour: str) -> None:
    spPr = el.find(qn("p:spPr"))
    if spPr is None:
        raise ValueError("this shape has no fill to change")
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        for old in spPr.findall(qn(tag)):
            spPr.remove(old)
    fill = etree.Element(qn("a:solidFill"))
    fill.append(etree.fromstring(token_xml(colour)))
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = spPr.find(qn("a:custGeom"))
    (geom.addnext(fill) if geom is not None else spPr.insert(0, fill))


def _set_text_colour(el, colour: str) -> None:
    for rpr in el.iter(qn("a:rPr"), qn("a:endParaRPr")):
        for old in rpr.findall(qn("a:solidFill")):
            rpr.remove(old)
        fill = etree.Element(qn("a:solidFill"))
        fill.append(etree.fromstring(token_xml(colour)))
        ln = rpr.find(qn("a:ln"))
        (ln.addnext(fill) if ln is not None else rpr.insert(0, fill))  # schema order: ln, fill, ...


def _scale_fonts(el, scale: float) -> int:
    """Scale explicit font sizes inside an inserted component; returns how many hit the floor."""
    clamped = 0
    floor = MIN_SCALED_PT * 100
    for node in el.iter(qn("a:rPr"), qn("a:endParaRPr"), qn("a:defRPr")):
        sz = node.get("sz")
        if sz and sz.isdigit():
            new = int(round(int(sz) * scale))
            if new < floor:
                new, clamped = floor, clamped + 1
            node.set("sz", str(new))
    return clamped


# --------------------------------------------------------------------------
# Text, preserving the template's formatting
# --------------------------------------------------------------------------
_BOLD = re.compile(r"(\*\*.+?\*\*)")


def _template_paragraph(txBody):
    ps = txBody.findall(qn("a:p"))
    for p in ps:
        if p.find(qn("a:r")) is not None:
            return p
    return ps[0] if ps else etree.SubElement(txBody, qn("a:p"))


def _normalise_items(value) -> list[dict]:
    if isinstance(value, str):
        value = value.split("\n")
    items = []
    for v in value:
        if isinstance(v, dict):
            items.append({"text": str(v.get("text", "")), "level": v.get("level")})
        else:
            items.append({"text": str(v), "level": None})
    return items


def set_text_frame(tf, value) -> list[str]:
    """Replace the text of a text frame, cloning the first paragraph/run formatting.

    Returns the plain paragraphs (for fit estimation).
    """
    txBody = tf._txBody
    template = _template_paragraph(txBody)
    base_rpr = None
    first_run = template.find(qn("a:r"))
    if first_run is not None and first_run.find(qn("a:rPr")) is not None:
        base_rpr = deepcopy(first_run.find(qn("a:rPr")))
    elif template.find(qn("a:endParaRPr")) is not None:
        base_rpr = deepcopy(template.find(qn("a:endParaRPr")))
        base_rpr.tag = qn("a:rPr")

    skeleton = deepcopy(template)
    for child in list(skeleton):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            skeleton.remove(child)

    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)

    plain = []
    for item in _normalise_items(value):
        p = deepcopy(skeleton)
        if item["level"] is not None:
            pPr = p.find(qn("a:pPr"))
            if pPr is None:
                pPr = etree.Element(qn("a:pPr"))
                p.insert(0, pPr)
            pPr.set("lvl", str(int(item["level"])))
        end = p.find(qn("a:endParaRPr"))
        for seg in _BOLD.split(item["text"]):
            if not seg:
                continue
            bold = seg.startswith("**") and seg.endswith("**") and len(seg) > 4
            text = seg[2:-2] if bold else seg
            r = etree.Element(qn("a:r"))
            if base_rpr is not None:
                rpr = deepcopy(base_rpr)
                if bold:
                    rpr.set("b", "1")
                r.append(rpr)
            elif bold:
                r.append(etree.Element(qn("a:rPr"), b="1"))
            t = etree.SubElement(r, qn("a:t"))
            t.text = text
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        txBody.append(p)
        plain.append(item["text"])
    return plain


# --------------------------------------------------------------------------
# Picture geometry
# --------------------------------------------------------------------------
def _apply_crop(pic, iw, ih, fw, fh, fit):
    if fit == "contain" or not (iw and ih and fw and fh):
        pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
        return
    frame_ar, img_ar = fw / fh, iw / ih
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
    if img_ar > frame_ar:
        c = (1 - frame_ar / img_ar) / 2
        pic.crop_left = pic.crop_right = c
    elif img_ar < frame_ar:
        c = (1 - img_ar / frame_ar) / 2
        pic.crop_top = pic.crop_bottom = c


def _align_in_slot(el, slot, slide_w) -> None:
    """Move a contained picture or shape group to the side of its slot: a logo in a left
    corner sits on the left margin, in a right corner on the right one, centred otherwise."""
    xfrm = el.find(".//" + qn("a:xfrm"))
    off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
    left, _, width, _ = slot
    w = int(ext.get("cx"))
    centre = left + width / 2
    if centre < slide_w / 3:
        x = left
    elif centre > slide_w * 2 / 3:
        x = left + width - w
    else:
        x = left + (width - w) // 2
    off.set("x", str(int(x)))


def _contain(pic, left, top, fw, fh, iw, ih):
    scale = min(fw / iw, fh / ih)
    w, h = int(iw * scale), int(ih * scale)
    pic.left, pic.top = int(left + (fw - w) / 2), int(top + (fh - h) / 2)
    pic.width, pic.height = w, h


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------
def write_credits(out: Path, credits: list[dict]) -> str | None:
    if not credits:
        return None
    path = out.with_suffix(".credits.md")
    lines = ["# Image credits", ""]
    for c in credits:
        lines.append(f"- Slide {c['slide']}: {c['credit']}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return str(path)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", required=True, help="deck spec JSON")
    ap.add_argument("--out", help="output .pptx (overrides spec.output)")
    ap.add_argument("--local-dir", help="folder with your own library files (see config.json for the names)")
    ap.add_argument("--layouts", "--design", dest="layouts", help="layouts library .pptx (skips the lookup)")
    ap.add_argument("--visuals", "--shapes", dest="visuals", help="visuals library .pptx (skips the lookup)")
    ap.add_argument("--layouts-index", "--design-index", dest="layouts_index", help="layouts index (.csv or .xlsx)")
    ap.add_argument("--visuals-index", "--shapes-index", dest="visuals_index", help="visuals index (.csv or .xlsx)")
    ap.add_argument("--strict", action="store_true", help="fail instead of warn on unknown shape names")
    args = ap.parse_args(argv)

    try:
        spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
        if not isinstance(spec.get("slides"), list) or not spec["slides"]:
            raise SpecError("spec needs a non-empty 'slides' list")
        spec_dir = Path(args.spec).resolve().parent
        logo = load_logo(spec, spec_dir)  # read it first: a wrong file stops the build before any work

        if args.layouts:
            paths = {"layouts": {"pptx": args.layouts, "index": args.layouts_index},
                     "visuals": {"pptx": args.visuals, "index": args.visuals_index}}
        else:
            t = ensure_templates(args.local_dir)
            paths = {"layouts": t["layouts"], "visuals": t["visuals"]}

        def _idx(p):
            try:
                return load_index(p) if p else []
            except Exception as exc:  # noqa: BLE001
                eprint(f"warning: index {p} not loaded ({exc})")
                return []

        out = Path(args.out or spec.get("output") or "deck.pptx")
        libs = {Path(p).resolve() for kind in paths.values() for p in [kind.get("pptx")] if p}
        if out.resolve() in libs:
            raise SpecError(f"output {out} would overwrite a library file; choose another path")

        b = Builder(
            paths["layouts"]["pptx"],
            paths["visuals"].get("pptx"),
            _idx(paths["layouts"].get("index")),
            _idx(paths["visuals"].get("index")),
            strict=args.strict,
            logo=logo,
            spec_dir=spec_dir,
        )
        b.warnings.extend(check_keys(spec))
        if spec.get("logo_alt") and not logo:
            b.warn("spec: 'logo_alt' is ignored without a 'logo'")
        for i, s in enumerate(spec["slides"]):
            if s.get("hidden"):
                continue
            b.add_slide(s, i)
        b.finish(spec.get("properties"), spec.get("language"))

        out.parent.mkdir(parents=True, exist_ok=True)
        b.prs.save(out)
        credits = write_credits(out, b.credits)
        print(json.dumps({
            "output": str(out),
            "slides": len(b.prs.slides),
            "warnings": b.warnings,
            "credits_file": credits,
        }, indent=2, ensure_ascii=False))
        return 0
    except (SpecError, FetchError, FileNotFoundError, ValueError) as exc:
        print(json.dumps({"error": str(exc)}, indent=2, ensure_ascii=False))
        return 2


if __name__ == "__main__":
    for _stream in (sys.stdout, sys.stderr):  # Windows pipes default to a legacy code page; agents read UTF-8
        getattr(_stream, "reconfigure", lambda **_: None)(encoding="utf-8")
    sys.exit(main())
