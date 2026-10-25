"""Slide and component cloning across presentations.

python-pptx has no public "copy slide" API. The reliable recipe is:

1. create a new slide on a matching layout,
2. deep-copy the source shape tree into it,
3. for every relationship the copied XML points at (pictures, charts, embedded
   workbooks, SmartArt...) copy the target part into the destination package and
   rewrite the r:id / r:embed / r:link attributes.

Media (images, video, audio) are shared when source and destination are the same
package, everything else (charts, workbooks) is always duplicated so editing one
cloned chart never changes another slide.
"""
from __future__ import annotations

import re
from copy import deepcopy

from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn

R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
R_ATTR_PREFIX = "{%s}" % R_NS

_SKIP_RELTYPES = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE, RT.SLIDE, RT.SLIDE_MASTER}
_SHAREABLE_PREFIXES = ("image/", "video/", "audio/", "font/")
_RID_IN_BLOB = re.compile(rb'(r:(?:id|embed|link|pict))="(rId\d+)"')


def _partname_template(partname: str) -> str:
    m = re.match(r"^(.*?)(\d+)(\.[^./]+)$", str(partname))
    if m:
        return m.group(1) + "%d" + m.group(3)
    base, dot, ext = str(partname).rpartition(".")
    return f"{base}%d{dot}{ext}" if dot else f"{partname}%d"


class SlideCloner:
    """Copies slides/shapes into `dst_prs`."""

    def __init__(self, dst_prs):
        self.dst_prs = dst_prs
        self.dst_package = dst_prs.part.package
        self._existing = {str(p.partname) for p in self.dst_package.iter_parts()}
        self._cache: dict[int, Part] = {}        # per transplant: charts, workbooks... always fresh copies
        self._media_cache: dict[int, Part] = {}  # whole build: an image used twice is stored once
        self.warnings: list[str] = []

    # ------------------------------------------------------------------ parts
    def _alloc_partname(self, src_partname: str) -> PackURI:
        tmpl = _partname_template(src_partname)
        n = 1
        while (tmpl % n) in self._existing:
            n += 1
        name = tmpl % n
        self._existing.add(name)
        return PackURI(name)

    def _copy_part(self, src_part) -> Part:
        is_media = src_part.content_type.startswith(_SHAREABLE_PREFIXES)
        cache = self._media_cache if is_media else self._cache
        if id(src_part) in cache:
            return cache[id(src_part)]
        if is_media and src_part.package is self.dst_package:
            cache[id(src_part)] = src_part
            return src_part

        # Same class as the source (ChartPart, EmbeddedXlsxPart, ...) so python-pptx
        # keeps treating the copy as a chart / workbook after cloning.
        new = type(src_part).load(
            self._alloc_partname(str(src_part.partname)),
            src_part.content_type,
            self.dst_package,
            src_part.blob,
        )
        cache[id(src_part)] = new  # cache first: relationship graphs can loop

        rid_map: dict[str, str] = {}
        for rId, rel in sorted(src_part.rels.items(), key=lambda kv: kv[0]):
            if rel.is_external:
                rid_map[rId] = new.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            else:
                target = self._copy_part(rel.target_part)
                rid_map[rId] = new.rels.get_or_add(rel.reltype, target)

        if any(old != nw for old, nw in rid_map.items()):
            element = getattr(new, "_element", None)
            if element is not None:  # XML part: rewrite attributes in place
                for node in element.iter():
                    for attr, val in list(node.attrib.items()):
                        if attr.startswith(R_ATTR_PREFIX) and val in rid_map:
                            node.set(attr, rid_map[val])
            else:  # opaque part: rewrite the bytes
                def _swap(m):
                    old = m.group(2).decode()
                    return m.group(1) + b'="' + rid_map.get(old, old).encode() + b'"'

                new._blob = _RID_IN_BLOB.sub(_swap, new.blob)
        return new

    # ----------------------------------------------------------- relationships
    def transplant(self, src_part, dst_part, elements) -> None:
        """Make `elements` (already deep-copied into dst) valid inside dst_part."""
        self._cache = {}  # each insertion gets its own chart/workbook copies
        rid_map: dict[str, str | None] = {}
        dangling = []
        for el in elements:
            for node in el.iter():
                for attr, val in list(node.attrib.items()):
                    if not attr.startswith(R_ATTR_PREFIX):
                        continue
                    if val not in rid_map:
                        rid_map[val] = self._map_rel(src_part, dst_part, val)
                    new_id = rid_map[val]
                    if new_id is None:
                        dangling.append(node)
                    else:
                        node.set(attr, new_id)
        for node in dangling:
            parent = node.getparent()
            if parent is not None:
                parent.remove(node)

    def _map_rel(self, src_part, dst_part, rId):
        rel = src_part.rels.get(rId)
        if rel is None:
            return None
        if rel.is_external:
            return dst_part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        if rel.reltype in _SKIP_RELTYPES:
            return None
        return dst_part.rels.get_or_add(rel.reltype, self._copy_part(rel.target_part))

    # ------------------------------------------------------------------ layouts
    def resolve_layout(self, src_slide):
        src_layout = src_slide.slide_layout
        if src_layout.part.package is self.dst_package:
            return src_layout
        by_name = {lay.name: lay for lay in self.dst_prs.slide_layouts}
        if src_layout.name in by_name:
            return by_name[src_layout.name]
        blank = by_name.get("Blank")
        if blank is None:
            blank = min(self.dst_prs.slide_layouts, key=lambda lay: len(lay.placeholders))
        self.warnings.append(
            f"Layout '{src_layout.name}' not found in the design template; "
            f"using '{blank.name}'. Keep both templates on the same master to avoid this."
        )
        return blank

    # ------------------------------------------------------------------- slides
    def clone_slide(self, src_slide):
        """Append a copy of `src_slide` to the destination deck and return it."""
        new = self.dst_prs.slides.add_slide(self.resolve_layout(src_slide))
        spTree = new.shapes._spTree
        keep = {qn("p:nvGrpSpPr"), qn("p:grpSpPr")}
        for el in list(spTree):
            if el.tag not in keep:
                spTree.remove(el)

        copied = []
        for el in src_slide.shapes._spTree:
            if el.tag in keep or el.tag == qn("p:extLst"):
                continue
            clone = deepcopy(el)
            spTree.append(clone)
            copied.append(clone)

        # slide background
        src_bg = src_slide._element.cSld.find(qn("p:bg"))
        if src_bg is not None:
            new_cSld = new._element.cSld
            old_bg = new_cSld.find(qn("p:bg"))
            if old_bg is not None:
                new_cSld.remove(old_bg)
            bg = deepcopy(src_bg)
            new_cSld.insert(0, bg)
            copied.append(bg)

        self.transplant(src_slide.part, new.part, copied)
        return new

    def clone_shape(self, src_slide, src_shape, dst_slide):
        """Copy one top-level shape (picture, chart, group...) into `dst_slide`."""
        clone = deepcopy(src_shape._element)
        dst_slide.shapes._spTree.append(clone)
        self.transplant(src_slide.part, dst_slide.part, [clone])
        for shp in dst_slide.shapes:
            if shp._element is clone:
                return shp
        raise RuntimeError("cloned shape not found on destination slide")


def delete_slide(prs, index: int) -> None:
    sldIdLst = prs.slides._sldIdLst
    sldId = sldIdLst[index]
    prs.part.drop_rel(sldId.rId)
    sldIdLst.remove(sldId)
