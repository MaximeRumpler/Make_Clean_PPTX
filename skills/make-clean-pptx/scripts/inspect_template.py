#!/usr/bin/env python3
"""Inspect a template deck, and bootstrap its index spreadsheet.

For people *maintaining* the templates (and for agents debugging a wrong key):

    inspect_template.py design.pptx                    # every slide, every named shape
    inspect_template.py design.pptx --slide 12         # one slide
    inspect_template.py layouts.pptx --emit-index layouts.csv [--kind layouts|visuals]   (.csv or .xlsx)
    inspect_template.py design.pptx --check            # naming-convention report

--emit-index writes a *draft*: layout name, shape names, types and estimated
character limits are filled automatically; category, description, use_when and
tags are left for a human (or an LLM that has looked at the rendering) to complete.
"""
from __future__ import annotations

import argparse
import re
import sys

from pptx import Presentation

from _common import estimate_fit, has_text, iter_shapes, shape_kind

GENERIC = re.compile(r"^(TextBox|Rectangle|Rounded Rectangle|Oval|Freeform|Group|Picture|Content Placeholder|"
                     r"Text Placeholder|Title|Chart|Table|Straight Connector|Isosceles Triangle|Shape|Google Shape)[ ;_-]*\d+$", re.I)
# PowerPoint's default names always end with a number ("TextBox 12", "Title 1"); "Title" alone is a deliberate name.


def slide_fields(slide) -> list[dict]:
    fields = []
    for shp in iter_shapes(slide):
        kind = shape_kind(shp)
        if kind == "group":
            continue
        if kind == "text" and not has_text(shp) and not getattr(shp, "is_placeholder", False):
            continue  # decorative empty shapes
        if kind == "other":
            continue
        entry = {"name": shp.name, "type": {"text": "text", "picture": "image", "chart": "chart", "table": "table"}.get(kind, kind)}
        if kind == "text":
            fit = estimate_fit(shp, ["x"])
            if fit:
                cap = fit["chars_per_line"] * fit["avail_lines"]
                entry["limit"] = int(cap * 0.85)
                entry["lines"] = fit["avail_lines"]
                if fit["avail_lines"] > 2 and shp.text_frame.text.count("\n") >= 1:
                    entry["type"] = "bullets"
            entry["sample"] = shp.text_frame.text.strip().replace("\n", " / ")[:50]
        elif kind == "picture" and shp.width and shp.height:
            w, h = int(shp.width), int(shp.height)
            ratio = w / h
            entry["hint"] = f"{ratio:.2f}:1"
        fields.append(entry)
    return fields


def fields_string(fields: list[dict]) -> str:
    out = []
    for f in fields:
        third = f.get("limit") or f.get("hint") or ""
        out.append(f"{f['name']}|{f['type']}|{third}".rstrip("|"))
    return "; ".join(out)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--slide", type=int)
    ap.add_argument("--emit-index", metavar="CSV_OR_XLSX")
    ap.add_argument("--kind", choices=["layouts", "visuals", "design", "shapes"], default="layouts")
    ap.add_argument("--check", action="store_true", help="report shapes whose names are generic/duplicated")
    args = ap.parse_args(argv)

    prs = Presentation(args.pptx)
    slides = list(enumerate(prs.slides, 1))
    if args.slide:
        slides = [s for s in slides if s[0] == args.slide]
        if not slides:
            print(f"no slide {args.slide} (deck has {len(prs.slides)})", file=sys.stderr)
            return 2

    if args.check:
        bad = 0
        for n, s in slides:
            names = [f["name"] for f in slide_fields(s)]
            dup = {x for x in names if names.count(x) > 1}
            gen = [x for x in names if GENERIC.match(x)]
            if dup or gen:
                bad += 1
                print(f"slide {n}: " + ("duplicate: " + ", ".join(sorted(dup)) + "  " if dup else "") +
                      ("generic: " + ", ".join(gen) if gen else ""))
        print(f"{bad} of {len(slides)} slides need clearer shape names" if bad else "all shape names look usable")
        return 1 if bad else 0

    visual = args.kind in ("visuals", "shapes")
    if args.emit_index and args.emit_index.lower().endswith(".csv"):
        import csv

        header = ["id", "slide", "name", "category", "kind" if visual else "style", "description", "use_when", "tags", "fields"]
        with open(args.emit_index, "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(header)
            for n, s in slides:
                w.writerow([f"{'V' if visual else 'L'}-{n:02d}", n, s.slide_layout.name, "", "", "", "", "",
                            fields_string(slide_fields(s))])
        print(f"wrote draft index with {len(slides)} rows to {args.emit_index}. Fill category/description/use_when/tags.")
        return 0

    if args.emit_index:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "index"
        prefix = "V" if visual else "L"
        header = ["slide", "id", "category", "name", "description", "use_when", "fields", "tags", "notes"]
        if visual:
            header.insert(4, "kind")
        ws.append(header)
        for n, s in slides:
            fields = slide_fields(s)
            row = {"slide": n, "id": f"{prefix}-{n:03d}", "category": "", "name": s.slide_layout.name,
                   "description": "", "use_when": "", "fields": fields_string(fields), "tags": "", "notes": ""}
            row["kind"] = ""
            ws.append([row[h] for h in header])
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1C2333")
        widths = {"slide": 7, "id": 9, "category": 16, "name": 26, "kind": 12, "description": 44, "use_when": 36, "fields": 70, "tags": 26, "notes": 30}
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 20)
            for cell in ws[get_column_letter(i)][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        wb.save(args.emit_index)
        print(f"wrote draft index with {len(slides)} rows to {args.emit_index}. Fill category/description/use_when/tags.")
        return 0

    for n, s in slides:
        print(f"--- slide {n}  layout: {s.slide_layout.name}")
        for f in slide_fields(s):
            extra = f"  ≤{f['limit']} chars (~{f['lines']} lines)" if "limit" in f else (f"  {f['hint']}" if "hint" in f else "")
            sample = f"  e.g. “{f['sample']}”" if f.get("sample") else ""
            print(f"    {f['name']!r:<34} {f['type']:<8}{extra}{sample}")
        comps = [sh.name for sh in s.shapes]
        print(f"    top-level shapes (usable as 'component'): {', '.join(map(repr, comps))}")
    return 0


if __name__ == "__main__":
    for _stream in (sys.stdout, sys.stderr):  # Windows pipes default to a legacy code page; agents read UTF-8
        getattr(_stream, "reconfigure", lambda **_: None)(encoding="utf-8")
    sys.exit(main())
