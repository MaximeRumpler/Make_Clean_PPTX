#!/usr/bin/env python3
"""Read and search the library indexes (layouts.csv, visuals.csv, or .xlsx).

The index is the catalogue of a library: one row per layout slide, or per
visual component. Reading it lets the agent choose without opening every slide.

Examples
--------
  read_index.py layouts --summary
  read_index.py layouts --query "three options decision" --top 6
  read_index.py visuals --query "variance budget actual" --where objective="Break down"
  read_index.py visuals --where frame=wide --where kind=chart
  read_index.py visuals --id V-COM-01
  read_index.py path/to/index.csv --format json

"layouts" and "visuals" resolve to the bundled indexes (or to your own, see
fetch_templates.py). Header names are matched loosely (see ALIASES); extra
columns are kept and can be used with --where. Contract: references/index-schema.md.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from pathlib import Path

ALIASES = {
    "slide": ["slide", "slide_no", "slide_number", "slide_nr", "no", "n", "#", "number", "page", "position"],
    "id": ["id", "slide_id", "uid", "key", "ref", "code"],
    "category": ["category", "section", "group", "family", "type"],
    "name": ["name", "title", "layout", "layout_name", "component", "component_name"],
    "description": ["description", "desc", "summary", "what"],
    "use_when": ["use_when", "usage", "use_case", "use_cases", "when", "best_for", "when_to_use"],
    "fields": ["fields", "placeholders", "slots", "editable", "editable_fields", "shapes"],
    "kind": ["kind", "element", "format", "component_type"],
    "tags": ["tags", "keywords", "labels"],
    "notes": ["notes", "remarks", "comments", "rules"],
    "objective": ["objective", "goal", "communication_goal", "intent"],
    "data_type": ["data_type", "data", "data_shape"],
}
TEXT_KEYS = ("id", "category", "name", "description", "use_when", "fields", "kind", "tags", "notes", "objective", "data_type")

_ALIAS_LOOKUP = {a: canon for canon, alist in ALIASES.items() for a in alist}


def _norm_header(value) -> str:
    return re.sub(r"[^a-z0-9#]+", "_", str(value or "").strip().lower()).strip("_")


def parse_fields(raw: str | None) -> list[dict]:
    """Parse 'Title|text|60; Number 1|text|4|optional,default,follows=Item 1; Photo|image|1.78:1' into dicts.

    Flags: optional (deleted when not filled), default (the template text is an
    acceptable value and is kept), follows=X (removed when X is removed).
    """
    out = []
    for chunk in re.split(r"[;\n]+", raw or ""):
        chunk = chunk.strip()
        if not chunk:
            continue
        parts = [p.strip() for p in chunk.split("|")]
        flags = [f.strip() for p in parts[3:] for f in p.split(",") if f.strip()]
        if len(parts) > 2 and parts[2] in ("optional", "default"):
            flags.append(parts[2])
        follows = next((f.split("=", 1)[1].strip() for f in flags if f.startswith("follows=")), None)
        entry = {"name": parts[0], "type": parts[1] if len(parts) > 1 and parts[1] else "text",
                 "limit": None, "hint": None, "optional": "optional" in flags, "default": "default" in flags,
                 "follows": follows}
        if len(parts) > 2 and parts[2] and parts[2] not in ("optional", "default"):
            if parts[2].isdigit():
                entry["limit"] = int(parts[2])
            else:
                entry["hint"] = parts[2]
        out.append(entry)
    return out


def resolve_path(name: str) -> str:
    """'layouts' / 'visuals' (or 'design' / 'shapes') -> the configured index file."""
    if name in ("layouts", "visuals", "design", "shapes"):
        from fetch_templates import ensure_templates

        return ensure_templates()[name]["index"]
    return name


def _rows(path: Path) -> list:
    if path.suffix.lower() in (".csv", ".tsv", ".txt"):
        text = path.read_text(encoding="utf-8-sig")
        delim = "\t" if path.suffix.lower() == ".tsv" else (";" if text.split("\n", 1)[0].count(";") >= text.split("\n", 1)[0].count(",") else ",")
        return [tuple(r) for r in csv.reader(text.splitlines(), delimiter=delim)]
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb["index"] if "index" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    return rows


def load_index(path: str | Path) -> list[dict]:
    """Return a list of normalised entries, in file order."""
    path = Path(resolve_path(str(path)))
    rows = _rows(path)
    header_row, mapping = None, {}
    for i, row in enumerate(rows[:10]):
        heads = list(map(_norm_header, row))
        m = {j: _ALIAS_LOOKUP.get(h, h) for j, h in enumerate(heads) if h}
        if len({_ALIAS_LOOKUP[h] for h in heads if h in _ALIAS_LOOKUP}) >= 2 and ("slide" in m.values() or "name" in m.values()):
            header_row, mapping = i, m
            break
    if header_row is None:
        raise ValueError(
            f"{path}: could not find a header row. Expected columns such as "
            "slide, id, category, name, description, fields (see references/index-schema.md)."
        )

    entries = []
    for row in rows[header_row + 1:]:
        if not any(c not in (None, "") for c in row):
            continue
        e = {}
        for j, key in mapping.items():
            if j < len(row) and row[j] not in (None, ""):
                e.setdefault(key, row[j])
        if "slide" in e:
            try:
                e["slide"] = int(float(e["slide"]))
            except (TypeError, ValueError):
                continue
        else:
            e["slide"] = len(entries) + 1
        for k in TEXT_KEYS:
            e[k] = str(e[k]).strip() if k in e else ""
        e["field_list"] = parse_fields(e["fields"])
        entries.append(e)
    return entries


def resolve(entries: list[dict], *, slide: int | None = None, id_: str | None = None) -> dict | None:
    for e in entries:
        if id_ and e["id"].lower() == id_.lower():
            return e
        if slide is not None and e["slide"] == slide and not id_:
            return e
    return None


# ---------------------------------------------------------------- searching
def _tokens(text: str) -> set[str]:
    return {t[:-1] if len(t) > 3 and t.endswith("s") else t for t in re.findall(r"[a-z0-9]+", text.lower()) if len(t) > 1}


WEIGHTS = {"name": 3.0, "tags": 3.0, "category": 2.0, "kind": 2.0, "objective": 2.0, "data_type": 1.0,
           "description": 1.0, "use_when": 1.5, "fields": 0.5}


def score(entry: dict, query: str) -> float:
    q = _tokens(query)
    if not q:
        return 0.0
    return sum(w * len(q & _tokens(str(entry.get(key, "")))) for key, w in WEIGHTS.items())


def search(entries: list[dict], query: str, top: int = 10) -> list[dict]:
    scored = [(score(e, query), e) for e in entries]
    scored = [x for x in scored if x[0] > 0]
    scored.sort(key=lambda x: (-x[0], x[1]["slide"]))
    return [e for _, e in scored[:top]]


def where(entries: list[dict], conditions: list[str]) -> list[dict]:
    out = entries
    for cond in conditions or []:
        key, _, val = cond.partition("=")
        key = _ALIAS_LOOKUP.get(_norm_header(key), _norm_header(key))
        val = val.strip().lower()
        out = [e for e in out if val in str(e.get(key, "")).lower()]
    return out


# ---------------------------------------------------------------- printing
SHOW = ("kind", "frame", "objective", "data_type", "style", "density", "description", "use_when", "limits", "fill_notes", "tags", "related", "notes")


def _fields_txt(e):
    return "; ".join(
        f"{f['name']}({f['type']}{',≤' + str(f['limit']) if f['limit'] else ''}{',' + f['hint'] if f['hint'] else ''}{',opt' if f['optional'] else ''})"
        for f in e["field_list"]
    )


def line(e: dict, detail: bool = False) -> str:
    head = f"#{e['slide']:<4} {e['id'] or '-':<9} [{e['category'] or '-'}] {e['name']}"
    if not detail:
        extra = " · ".join(str(e[k]) for k in ("objective", "frame") if e.get(k))
        return f"{head}{'  (' + extra + ')' if extra else ''} :: {e['description']}"
    parts = [head]
    for k in SHOW:
        if e.get(k):
            parts.append(f"    {k}: {e[k]}")
    if e["field_list"]:
        parts.append(f"    fields: {_fields_txt(e)}")
    return "\n".join(parts)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("index", help="'layouts', 'visuals', or a path to a .csv / .xlsx index")
    ap.add_argument("--summary", action="store_true", help="count entries per category")
    ap.add_argument("--query", "-q", help="free-text search")
    ap.add_argument("--category", "-c", help="filter by category (case-insensitive substring)")
    ap.add_argument("--where", "-w", action="append", help="filter key=value (substring), repeatable")
    ap.add_argument("--slide", type=int, help="show one entry by slide number")
    ap.add_argument("--id", help="show one entry by id")
    ap.add_argument("--top", type=int, default=10)
    ap.add_argument("--detail", action="store_true")
    ap.add_argument("--format", choices=["text", "json"], default="text")
    args = ap.parse_args(argv)

    try:
        entries = load_index(args.index)
    except Exception as exc:  # noqa: BLE001
        print(f"error: {exc}", file=sys.stderr)
        return 2

    if args.summary:
        counts = Counter(e["category"] or "(none)" for e in entries)
        if args.format == "json":
            print(json.dumps({"total": len(entries), "categories": counts}, indent=2))
        else:
            print(f"{len(entries)} entries")
            for cat, n in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0])):
                print(f"  {n:>3}  {cat}")
        return 0

    selected = entries
    if args.slide is not None or args.id:
        hit = resolve(entries, slide=args.slide, id_=args.id)
        selected = [hit] if hit else []
        args.detail = True
    else:
        if args.category:
            selected = [e for e in selected if args.category.lower() in e["category"].lower()]
        selected = where(selected, args.where)
        if args.query:
            selected = search(selected, args.query, args.top)

    if not selected:
        print("no match", file=sys.stderr)
        return 1
    if args.format == "json":
        print(json.dumps(selected, indent=2, ensure_ascii=False))
    else:
        for e in selected:
            print(line(e, args.detail))
    return 0


if __name__ == "__main__":
    for _stream in (sys.stdout, sys.stderr):  # Windows pipes default to a legacy code page; agents read UTF-8
        getattr(_stream, "reconfigure", lambda **_: None)(encoding="utf-8")
    sys.exit(main())
