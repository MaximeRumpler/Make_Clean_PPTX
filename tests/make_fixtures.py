"""Generate a tiny layouts library + visuals library (with indexes) for the tests.

The real libraries live in skills/make-clean-pptx/assets/. These fixtures follow
the same contract (named shapes, one index row per slide) so the pipeline can be
exercised in isolation: python tests/make_fixtures.py <out_dir>
Indexes are written as CSV (the default format) and as xlsx (still supported).
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

W, H = 12192000, 6858000  # 16:9
INK = RGBColor(0x1C, 0x23, 0x33)
ACCENT = RGBColor(0x10, 0xB9, 0x81)


def _new_prs() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    return prs


def _blank(prs):
    layout = [lay for lay in prs.slide_layouts if lay.name == "Blank"][0]
    return prs.slides.add_slide(layout)


def _text(slide, name, x, y, w, h, text, size=18, bold=False, color=INK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = text
    r = tf.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    return tb


def _png(path: Path, size, color):
    im = Image.new("RGB", size, color)
    ImageDraw.Draw(im).ellipse([size[0] * .2, size[1] * .2, size[0] * .8, size[1] * .8], fill="white")
    im.save(path)


def make(out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)

    # ---------------------------------------------------------------- design
    prs = _new_prs()
    s = _blank(prs)
    _text(s, "Title", 0.8, 2.4, 11.5, 1.4, "Cover title", 44, True)
    _text(s, "Subtitle", 0.8, 3.9, 11.5, 0.8, "Cover subtitle", 22)

    s = _blank(prs)
    _text(s, "Title", 0.8, 0.5, 11.5, 1.0, "Bullets title", 32, True)
    body = _text(s, "Body", 0.8, 1.8, 11.5, 4.5, "First bullet", 20)
    body.text_frame.add_paragraph().text = "Second bullet"

    s = _blank(prs)
    _text(s, "Title", 0.8, 0.5, 11.5, 1.0, "Image slide title", 32, True)
    _text(s, "Body", 0.8, 1.8, 5.5, 4.0, "Body text", 18)
    img = out / "_placeholder.png"
    _png(img, (800, 800), (200, 200, 200))
    pic = s.shapes.add_picture(str(img), Inches(7), Inches(1.5), Inches(5.5), Inches(5))
    pic.name = "Photo"

    s = _blank(prs)
    _text(s, "Title", 0.8, 0.5, 11.5, 1.0, "Stats title", 32, True)
    for i in (1, 2, 3):
        _text(s, f"Number {i}", 0.8 + (i - 1) * 4, 2.5, 3.5, 1.2, "00%", 54, True, ACCENT)
        _text(s, f"Label {i}", 0.8 + (i - 1) * 4, 3.8, 3.5, 1.0, "Label", 16)

    s = _blank(prs)
    _text(s, "Title", 0.8, 0.5, 11.5, 1.0, "Table title", 32, True)
    gf = s.shapes.add_table(3, 2, Inches(0.8), Inches(1.8), Inches(11), Inches(2.4))
    gf.name = "Table"
    for r in range(3):
        for c in range(2):
            gf.table.cell(r, c).text = f"r{r}c{c}"

    s = _blank(prs)
    _text(s, "Title", 0.8, 0.5, 11.5, 1.0, "Frame title", 32, True)
    frame = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5))
    frame.name = "Content Area"

    design = out / "layouts.pptx"
    prs.save(design)

    # ---------------------------------------------------------------- shapes
    prs2 = _new_prs()
    s = _blank(prs2)
    grp = s.shapes.add_group_shape()
    grp.name = "Timeline 4 steps"
    for i in range(4):
        box = grp.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5 + i * 3), Inches(2), Inches(2.6), Inches(1.6))
        box.name = f"Step {i + 1} title"
        box.text_frame.text = f"Step {i + 1}"
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
        dot = grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.5 + i * 3), Inches(4), Inches(0.4), Inches(0.4))
        dot.name = f"Dot {i + 1}"

    s = _blank(prs2)
    cd = CategoryChartData()
    cd.categories = ["A", "B", "C"]
    cd.add_series("Old", (1, 2, 3))
    gfc = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(8), Inches(5), cd)
    gfc.name = "Bar Chart"

    shapes = out / "visuals.pptx"
    prs2.save(shapes)

    # ---------------------------------------------------------------- indexes
    def index(path, rows, kind=False):
        header = ["Slide", "ID", "Category", "Name", "Description", "Use when", "Fields", "Tags"]
        wb = Workbook()
        ws = wb.active
        ws.title = "index"
        ws.append(header)
        for r in rows:
            ws.append(r)
        wb.save(path)
        with open(path.with_suffix(".csv").with_name(path.stem.replace("_index", "") + ".csv"), "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(header)
            w.writerows(rows)

    index(out / "layouts_index.xlsx", [
        [1, "D-001", "Cover", "Cover", "Opening slide with big title", "start of deck", "Title|text|60; Subtitle|text|90", "opener title"],
        [2, "D-002", "Content", "Bullets", "Title and bullet list", "explain several points", "Title|text|70; Body|bullets|400", "list bullets"],
        [3, "D-003", "Content", "Image right", "Text left, photo right", "illustrate a message", "Title|text|70; Body|text|250; Photo|image|1.10:1", "photo image"],
        [4, "D-004", "Data", "Three stats", "Three big numbers", "key figures kpi", "Title|text|70; Number 1|text|6; Label 1|text|40", "numbers kpi stats"],
        [5, "D-005", "Data", "Table", "Simple table", "compare values", "Title|text|70; Table|table", "table comparison"],
        [6, "D-006", "Frame", "Content frame", "Title and an empty area for a component", "diagrams charts", "Title|text|70; Content Area|frame", "diagram frame"],
    ])
    index(out / "visuals_index.xlsx", [
        [1, "S-001", "Process", "Timeline 4 steps", "Horizontal timeline", "roadmap phases", "Step 1 title|text|20", "timeline roadmap"],
        [2, "S-002", "Chart", "Bar Chart", "Clustered columns", "compare categories", "Bar Chart|chart", "chart bar"],
    ])
    return {"design": design, "shapes": shapes, "design_index": out / "layouts.csv",
            "shapes_index": out / "visuals.csv", "design_xlsx": out / "layouts_index.xlsx", "image": img}


if __name__ == "__main__":
    paths = make(Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures"))
    for k, v in paths.items():
        print(k, v)
