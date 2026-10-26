"""The bundled libraries, their indexes and the features built on them."""
from __future__ import annotations

import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / "skills" / "make-clean-pptx"
ASSETS = SKILL / "assets"
SCRIPTS = SKILL / "scripts"
sys.path.insert(0, str(SCRIPTS))

from _common import abs_box, iter_shapes  # noqa: E402
from _svg import parse_path, parse_svg, svg_to_shapes  # noqa: E402
from build_deck import waterfall_series  # noqa: E402
from read_index import load_index  # noqa: E402


def run(script, *args):
    p = subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)], capture_output=True, encoding="utf-8")
    return p.returncode, p.stdout, p.stderr


# ------------------------------------------------------------------ contract
@pytest.mark.parametrize("lib", ["layouts", "visuals"])
def test_index_matches_pptx(lib):
    prs = Presentation(ASSETS / f"{lib}.pptx")
    entries = load_index(ASSETS / f"{lib}.csv")
    assert len(entries) == len(prs.slides)
    ids = [e["id"] for e in entries]
    assert len(set(ids)) == len(ids)
    for e in entries:
        slide = prs.slides[e["slide"] - 1]
        if lib == "visuals":
            groups = [s for s in slide.shapes if s.name == e["id"]]
            assert len(groups) == 1, f"{e['id']}: one group named after the id"
            names = {s.name for s in iter_shapes(groups[0])}
        else:
            names = {s.name for s in iter_shapes(slide)}
        for f in e["field_list"]:
            assert f["name"] in names, f"{e['id']}: field {f['name']} missing in the pptx"
        for key in ("category", "description", "use_when", "tags"):
            assert e[key], f"{e['id']}: empty {key}"
        assert len(e["tags"].split()) >= 5


@pytest.mark.parametrize("lib", ["layouts", "visuals"])
def test_colours_are_theme_references(lib):
    with zipfile.ZipFile(ASSETS / f"{lib}.pptx") as z:
        for name in z.namelist():
            if re.match(r"ppt/(slides/slide|charts/chart)\d+\.xml$", name):
                xml = z.read(name).decode("utf-8")
                assert "srgbClr" not in xml, f"{name}: hard-coded colour (draw with the tokens)"
                assert "MCP " not in xml, f"{name}: sentinel font left"


def test_placeholders_are_not_content():
    """Template text describes the slot; it never contains names, dates or real figures."""
    pattern = re.compile(r"\b(19|20)\d\d\b|€|\$|lorem", re.I)
    for lib in ("layouts", "visuals"):
        for slide in Presentation(ASSETS / f"{lib}.pptx").slides:
            for shp in iter_shapes(slide):
                if getattr(shp, "has_text_frame", False):
                    assert not pattern.search(shp.text_frame.text), f"{lib}: '{shp.text_frame.text}'"


def test_icons_bundle():
    data = json.loads((ASSETS / "icons" / "lucide.json").read_text(encoding="utf-8"))
    assert len(data) > 1500 and "rocket" in data and data["rocket"]["tags"]
    assert (ASSETS / "icons" / "LICENSE-lucide.txt").is_file()


# ------------------------------------------------------------------ example
def test_example_deck_builds_clean(tmp_path):
    spec = json.loads((ROOT / "examples" / "home-energy.json").read_text(encoding="utf-8"))
    spec["output"] = str(tmp_path / "home.pptx")
    (tmp_path / "spec.json").write_text(json.dumps(spec))
    code, out, err = run("build_deck.py", "--spec", tmp_path / "spec.json")
    rep = json.loads(out)
    assert code == 0 and rep["warnings"] == [], rep
    code, out, _ = run("lint_deck.py", spec["output"], "--json")
    assert code == 0 and json.loads(out)["errors"] == []
    text = " ".join(s.text_frame.text for sl in Presentation(spec["output"]).slides for s in iter_shapes(sl)
                    if getattr(s, "has_text_frame", False))
    for ph in ("Slide title states the takeaway", "Short text", "Section", "Source or confidentiality note"):
        assert ph not in text, f"placeholder left: {ph}"


def _build(tmp_path, slides):
    spec = {"output": str(tmp_path / "t.pptx"), "slides": slides}
    (tmp_path / "s.json").write_text(json.dumps(spec))
    code, out, err = run("build_deck.py", "--spec", tmp_path / "s.json")
    assert code == 0, out + err
    return json.loads(out), Presentation(spec["output"])


def test_auto_clear_optional_and_followers(tmp_path):
    rep, prs = _build(tmp_path, [{"id": "L-06", "fill": {"Item 1": "First", "Item 2": "Second", "Item 3": None}}])
    names = {s.name for s in prs.slides[0].shapes}
    assert {"Item 1", "Number 1", "Item 2", "Number 2"} <= names
    assert "Number 3" not in names and "Item 3 mark" not in names   # follows its item
    assert "Item 6" not in names and "Kicker" not in names          # optional, not filled
    assert "Title" in names                                         # default text "Agenda" kept


def test_rich_values_position_and_shade(tmp_path):
    rep, prs = _build(tmp_path, [{"id": "L-31", "fill": {"Title": "Matrix", "Subtitle": None}, "insert": [
        {"id": "V-MAT-01", "component": "V-MAT-01", "into": "Content Area",
         "fill": {"Item 1": {"text": "1", "at": [50, 50], "fill": "a3"}, "Key 1": "One"}}]}])
    grp = [s for s in prs.slides[0].shapes if s.name == "V-MAT-01"][0]
    shapes = {s.name: s for s in iter_shapes(grp)}
    plot, item = abs_box(shapes["Plot Area"]), abs_box(shapes["Item 1"])
    centre = ((plot[0] + plot[2]) / 2, (plot[1] + plot[3]) / 2)
    assert abs((item[0] + item[2]) / 2 - centre[0]) < 20000 and abs((item[1] + item[3]) / 2 - centre[1]) < 20000
    assert 'val="accent3"' in etree.tostring(shapes["Item 1"]._element).decode()
    assert "Item 2" not in shapes and "Key 2 number" not in shapes


def test_icon_slot_becomes_native_shape(tmp_path):
    rep, prs = _build(tmp_path, [{"id": "L-20", "fill": {"Title": "Icons", "Heading 1": "A", "Body 1": "a",
                                                         "Icon 1": {"icon": "lucide:rocket"}, "Icon 2": None, "Icon 3": None,
                                                         "Heading 2": "B", "Body 2": "b", "Heading 3": "C", "Body 3": "c"}}])
    icon = [s for s in prs.slides[0].shapes if s.name == "Icon 1"][0]
    xml = etree.tostring(icon._element).decode()
    assert "custGeom" in xml and "cubicBezTo" in xml and "schemeClr" in xml


def test_waterfall_and_gantt_shortcuts(tmp_path):
    s = waterfall_series([100, 20, -30, 90])
    assert s == {"Base": [0, 100, 90, 0], "Total": [100, 0, 0, 90], "Increase": [0, 20, 0, 0], "Decrease": [0, 0, 30, 0]}
    rep, prs = _build(tmp_path, [{"id": "L-31", "fill": {"Title": "Plan", "Subtitle": None}, "insert": [
        {"id": "V-TIM-03", "component": "V-TIM-03", "into": "Content Area",
         "fill": {"Chart": {"categories": ["A", "B"], "gantt": [[0, 2], [1, 3]]}}}]}])
    chart = [s for s in iter_shapes(prs.slides[0]) if s.name == "Chart"][0].chart
    assert [list(p.values) for p in chart.plots[0].series] == [[0, 1], [2, 3]]


# ------------------------------------------------------------------ svg
def test_svg_path_packed_arc_flags_and_shapes():
    segs = parse_path("M12 5a.5.5 0 01.471.666L8 20a1 1 0 11-2 0z")
    assert segs[0] == ("M", 12.0, 5.0) and segs[-1] == ("Z",)
    assert any(s[0] == "C" for s in segs) and ("L", 8.0, 20.0) in segs
    vb, items, skipped = parse_svg('<svg viewBox="0 0 24 24" fill="none" stroke="currentColor"><circle cx="12" cy="12" r="5"/>'
                                   '<rect x="2" y="2" width="4" height="4" rx="1" fill="#000" fill-opacity=".2"/><text>x</text></svg>')
    assert vb == (0, 0, 24, 24) and len(items) == 2 and skipped == ["text"]
    xml, info = svg_to_shapes('<svg viewBox="0 0 24 24"><rect width="24" height="24" fill="#000" opacity="0.5"/></svg>', 0, 0, 914400, 914400)
    assert '<a:alpha val="50000"/>' in xml
    etree.fromstring(xml)  # well-formed


def test_find_icon_offline_search():
    code, out, _ = run("find_icon.py", "shield security", "--top", "3")
    res = json.loads(out)["results"]
    assert code == 0 and res and all(r["source"] == "lucide" for r in res)


def test_spec_typos_are_reported_and_libraries_protected(tmp_path):
    rep, prs = _build(tmp_path, [{"id": "L-12", "fil": {"Message": "x"}, "fill": {"Message": "Hello"},
                                  "insert": [], "notez": "n"}])
    assert any("unknown key 'fil'" in w and "did you mean 'fill'" in w for w in rep["warnings"])
    assert any("'notez'" in w for w in rep["warnings"])
    spec = {"output": str(ASSETS / "layouts.pptx"), "slides": [{"id": "L-12"}]}
    (tmp_path / "bad.json").write_text(json.dumps(spec))
    code, out, _ = run("build_deck.py", "--spec", tmp_path / "bad.json")
    assert code == 2 and "overwrite a library file" in out


def test_svg_input_hardening(tmp_path):
    import _svg
    with pytest.raises(_svg.SvgError, match="DOCTYPE"):
        parse_svg('<!DOCTYPE svg [<!ENTITY a "aaaa">]><svg viewBox="0 0 1 1"><path d="M0 0L1 1" stroke="#000"/></svg>')
    import find_icon
    (tmp_path / "notes.txt").write_text("secret")
    with pytest.raises(find_icon.IconError, match="only .svg"):
        find_icon.get_svg(f"local:{tmp_path / 'notes.txt'}")


def test_accessibility_language_and_decorative(tmp_path):
    with zipfile.ZipFile(ASSETS / "layouts.pptx") as z:
        assert "adec:decorative" in z.read("ppt/slides/slide6.xml").decode()  # agenda rules
    spec = {"output": str(tmp_path / "fr.pptx"), "language": "fr-FR", "slides": [{"id": "L-12", "fill": {"Message": "Bonjour"}}]}
    (tmp_path / "s.json").write_text(json.dumps(spec))
    code, out, _ = run("build_deck.py", "--spec", tmp_path / "s.json")
    assert code == 0, out
    xml = etree.tostring(Presentation(spec["output"]).slides[0]._element).decode()
    assert 'lang="fr-FR"' in xml and 'lang="en-US"' not in xml
