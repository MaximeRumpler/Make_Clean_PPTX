"""doctor.py --fix and render_check.py's retry, without touching the real Python or needing LibreOffice."""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "skills" / "make-clean-pptx" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import doctor  # noqa: E402
import render_check  # noqa: E402

ALL = {"python-pptx": "1", "openpyxl": "1", "Pillow": "1", "lxml": "1"}


# -------------------------------------------------------------------- doctor
def _doctor(capsys, *args):
    code = doctor.main(["--offline", *args])
    return code, json.loads(capsys.readouterr().out)


def test_doctor_reports_without_fixing(capsys, monkeypatch):
    monkeypatch.setattr(doctor, "check_packages", lambda: dict(ALL, openpyxl=None))
    monkeypatch.setattr(doctor, "pip_install", lambda user: pytest.fail("pip must not run without --fix"))
    code, rep = _doctor(capsys)
    assert code == 1 and rep["ok"] is False and "install" not in rep
    assert "openpyxl" in rep["message"] and "--fix" in rep["message"]


def test_doctor_fix_falls_back_to_user_install(capsys, monkeypatch):
    calls, state = [], {"present": False}

    def fake_pip(user):
        calls.append(user)
        state["present"] = user  # only the --user attempt succeeds
        return {"command": "pip", "ok": user, "error": None if user else "permission denied"}

    monkeypatch.setattr(doctor, "pip_install", fake_pip)
    monkeypatch.setattr(doctor, "check_packages", lambda: ALL if state["present"] else dict(ALL, Pillow=None))
    monkeypatch.setattr(sys, "base_prefix", sys.prefix)  # behave as outside a virtual environment
    code, rep = _doctor(capsys, "--fix")
    assert code == 0 and rep["ok"] is True
    assert calls == [False, True]
    assert rep["install"]["installed"] == ["Pillow"] and rep["install"]["still_missing"] == []
    assert "ready" in rep["message"]


def test_doctor_fix_reports_what_is_still_missing(capsys, monkeypatch):
    monkeypatch.setattr(doctor, "pip_install", lambda user: {"command": "pip", "ok": False, "error": "externally managed"})
    monkeypatch.setattr(doctor, "check_packages", lambda: dict(ALL, **{"python-pptx": None}))
    monkeypatch.setattr(sys, "base_prefix", sys.prefix)
    code, rep = _doctor(capsys, "--fix")
    assert code == 1 and len(rep["install"]["attempts"]) == 2
    assert rep["install"]["still_missing"] == ["python-pptx"]
    assert "python-pptx" in rep["message"] and "pip install" in rep["message"]


def test_pip_command_never_breaks_system_packages(monkeypatch):
    seen = []
    monkeypatch.setattr(doctor.subprocess, "run",
                        lambda cmd, **kw: seen.append(cmd) or subprocess.CompletedProcess(cmd, 1, "", "nope"))
    assert doctor.pip_install(user=False) == {"command": " ".join(seen[0]), "ok": False, "error": "nope"}
    doctor.pip_install(user=True)
    assert seen[1][-1] == "--user"
    assert not any("--break-system-packages" in c for cmd in seen for c in cmd)
    assert seen[0][:4] == [sys.executable, "-m", "pip", "install"]


def test_check_packages_sees_this_environment():
    assert all(doctor.check_packages().values())  # the test environment has every requirement


# -------------------------------------------------------------------- render
def _deck(path: Path) -> Path:
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))
    return path


def test_render_retries_on_a_resaved_copy(tmp_path, monkeypatch):
    """LibreOffice refuses the original (silently, as it does), accepts the re-saved copy."""
    deck = _deck(tmp_path / "deck.pptx")
    before = deck.read_bytes()
    seen = []

    def fake_run(cmd, **kw):
        if cmd[0] == "soffice":
            src, out = Path(cmd[-1]), Path(cmd[cmd.index("--outdir") + 1])
            seen.append(src)
            if src != deck:
                assert len(Presentation(str(src)).slides) == 1  # a real, readable copy
                (out / (src.stem + "." + cmd[cmd.index("--convert-to") + 1])).write_bytes(b"%PDF-1.4")
        else:  # pdftoppm
            Image.new("RGB", (64, 36), "white").save(cmd[-1] + "-1.png")
        return subprocess.CompletedProcess(cmd, 0)

    monkeypatch.setattr(render_check, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(render_check.shutil, "which", lambda name: "pdftoppm")
    monkeypatch.setattr(render_check.subprocess, "run", fake_run)
    info = render_check.render(deck, tmp_path / "r", dpi=40, cols=3, odp=True)
    assert info["resaved_copy"] is True and len(info["slides"]) == 1
    assert seen[0] == deck and seen[1] != deck and seen[1].name == deck.name
    assert Path(info["odp"]) == deck.with_suffix(".odp") and deck.with_suffix(".odp").is_file()
    assert deck.read_bytes() == before  # the user's file is never touched


def test_render_gives_up_after_one_retry(tmp_path, monkeypatch, capsys):
    deck = _deck(tmp_path / "deck.pptx")

    def refuse(cmd, **kw):
        raise subprocess.CalledProcessError(1, cmd)

    monkeypatch.setattr(render_check, "find_soffice", lambda: "soffice")
    monkeypatch.setattr(render_check.shutil, "which", lambda name: "pdftoppm")
    monkeypatch.setattr(render_check.subprocess, "run", refuse)
    assert render_check.main([str(deck), "--out-dir", str(tmp_path / "r")]) == 5
    rep = json.loads(capsys.readouterr().out)
    assert "re-saving" in rep["error"] and rep["hint"]


def test_render_reports_missing_tools(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(render_check, "find_soffice", lambda: None)
    assert render_check.main([str(_deck(tmp_path / "d.pptx"))]) == 4
    assert "LibreOffice" in json.loads(capsys.readouterr().out)["hint"]


@pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftoppm")), reason="LibreOffice/poppler not installed")
def test_render_retry_with_libreoffice(tmp_path, monkeypatch):
    """Real conversion; the first LibreOffice call is made to fail so the retry path runs."""
    deck = _deck(tmp_path / "deck.pptx")
    real, calls = subprocess.run, []

    def flaky(cmd, **kw):
        calls.append(cmd)
        if len(calls) == 1:
            raise subprocess.CalledProcessError(1, cmd, b"", b"Error: source file could not be loaded")
        return real(cmd, **kw)

    monkeypatch.setattr(render_check.subprocess, "run", flaky)
    info = render_check.render(deck, tmp_path / "r", dpi=30, cols=3)
    assert info["resaved_copy"] is True and len(info["slides"]) == 1
