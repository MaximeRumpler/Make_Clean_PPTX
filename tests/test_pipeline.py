"""End-to-end tests: build a deck from fixture templates and verify the result."""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "skills" / "make-clean-pptx" / "scripts"
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(ROOT / "tests"))

from make_fixtures import make  # noqa: E402
from read_index import load_index, parse_fields, search  # noqa: E402


@pytest.fixture(scope="session")
def fx(tmp_path_factory):
    return make(tmp_path_factory.mktemp("fixtures"))


def run(script, *args):
    p = subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)], capture_output=True, encoding="utf-8")
    return p.returncode, p.stdout, p.stderr


def build(fx, tmp_path, spec, extra=()):
    spec_path = tmp_path / "spec.json"
    spec.setdefault("output", str(tmp_path / "deck.pptx"))
    spec_path.write_text(json.dumps(spec))
    code, out, err = run("build_deck.py", "--spec", spec_path, "--design", fx["design"], "--shapes", fx["shapes"],
                         "--design-index", fx["design_index"], "--shapes-index", fx["shapes_index"], *extra)
    return code, json.loads(out), err


# ------------------------------------------------------------------ index
def test_index_parses_and_searches(fx):
    entries = load_index(fx["design_index"])
    assert len(entries) == 6
    assert entries[2]["field_list"][2] == {"name": "Photo", "type": "image", "limit": None, "hint": "1.10:1",
                                           "optional": False, "default": False, "follows": None}
    assert search(entries, "kpi numbers")[0]["slide"] == 4
    assert parse_fields("A|text|60; B|bullets|300")[1]["limit"] == 300


# ------------------------------------------------------------------ build
def test_basic_build_text_and_order(fx, tmp_path):
    code, rep, err = build(fx, tmp_path, {"slides": [
        {"slide": 1, "fill": {"Title": "Hello **bold**", "Subtitle": "World"}},
        {"slide": 2, "fill": {"Title": "Points", "Body": ["One", {"text": "Sub", "level": 1}, "Three"]}, "notes": "note!"},
        {"id": "D-001", "fill": {"Title": "Reused cover", "Subtitle": "again"}},
    ]})
    assert code == 0, rep
    assert rep["slides"] == 3 and rep["warnings"] == []
    prs = Presentation(rep["output"])
    assert len(prs.slides) == 3
    s1 = {s.name: s for s in prs.slides[0].shapes}
    assert s1["Title"].text_frame.text == "Hello bold"
    runs = s1["Title"].text_frame.paragraphs[0].runs
    assert [r.font.bold for r in runs] == [True, True] or runs[1].font.bold  # second run is bold
    assert runs[0].font.size.pt == 44  # template formatting preserved
    body = {s.name: s for s in prs.slides[1].shapes}["Body"].text_frame
    assert [p.text for p in body.paragraphs] == ["One", "Sub", "Three"]
    assert body.paragraphs[1].level == 1
    assert prs.slides[1].notes_slide.notes_text_frame.text == "note!"
    assert {s.name: s for s in prs.slides[2].shapes}["Title"].text_frame.text == "Reused cover"


def test_no_duplicate_zip_entries_and_templates_unchanged(fx, tmp_path):
    before = Path(fx["design"]).read_bytes()
    code, rep, _ = build(fx, tmp_path, {"slides": [{"slide": 3, "fill": {"Title": "x"}}, {"slide": 3, "fill": {"Title": "y"}}]})
    assert code == 0
    names = zipfile.ZipFile(rep["output"]).namelist()
    assert len(names) == len(set(names))
    assert Path(fx["design"]).read_bytes() == before
    # original 6 slides are gone, only the two clones remain
    assert len(Presentation(rep["output"]).slides) == 2


def test_image_replace_cover_crop_and_alt(fx, tmp_path):
    from PIL import Image
    wide = tmp_path / "wide.png"
    Image.new("RGB", (1600, 400), (10, 120, 200)).save(wide)
    code, rep, _ = build(fx, tmp_path, {"slides": [
        {"slide": 3, "fill": {"Title": "Img", "Body": "t", "Photo": {"path": str(wide), "alt": "A wide blue image", "credit": "Me, CC0"}}},
    ]})
    assert code == 0 and rep["warnings"] == []
    pic = {s.name: s for s in Presentation(rep["output"]).slides[0].shapes}["Photo"]
    assert pic.image.size == (1600, 400)
    assert pic.crop_left > 0 and pic.crop_right > 0 and pic.crop_top == 0  # wide image cropped sideways
    assert pic._element.xpath(".//p:cNvPr")[0].get("descr") == "A wide blue image"
    assert Path(rep["credits_file"]).read_text(encoding="utf-8").count("Me, CC0") == 1
    notes = Presentation(rep["output"]).slides[0].notes_slide.notes_text_frame.text
    assert "Me, CC0" in notes


def test_relative_input_paths_deck_folder_first_then_working_dir(fx, tmp_path):
    """Images, SVG files and local icons: folder of deck.json first, then the working directory."""
    from PIL import Image
    spec_dir, cwd = tmp_path / "spec", tmp_path / "cwd"
    (spec_dir / "img").mkdir(parents=True)
    (cwd / "img").mkdir(parents=True)
    Image.new("RGB", (400, 400), (10, 120, 200)).save(spec_dir / "img" / "a.png")   # next to deck.json
    Image.new("RGB", (400, 400), (200, 20, 20)).save(cwd / "img" / "b.png")        # only in the working dir
    Image.new("RGB", (400, 400), (0, 0, 0)).save(cwd / "img" / "a.png")            # same name: must lose
    svg = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24"><path d="M2 2h20v20H2z"/></svg>'
    (spec_dir / "img" / "logo.svg").write_text(svg, encoding="utf-8")
    absolute = tmp_path / "abs.png"
    Image.new("RGB", (400, 400), (0, 200, 0)).save(absolute)
    spec = {"output": str(tmp_path / "deck.pptx"), "slides": [
        {"slide": 3, "fill": {"Title": "a", "Body": "t", "Photo": {"path": "img/a.png", "alt": "x"}}},
        {"slide": 3, "fill": {"Title": "b", "Body": "t", "Photo": {"path": "img/b.png", "alt": "x"}}},
        {"slide": 3, "fill": {"Title": "abs", "Body": {"svg": "img/logo.svg"},
                              "Photo": {"path": str(absolute), "alt": "x"}}},
        {"slide": 3, "fill": {"Title": "icon", "Body": {"icon": "local:img/logo.svg"}, "Photo": None}},
    ]}
    (spec_dir / "deck.json").write_text(json.dumps(spec))
    p = subprocess.run([sys.executable, str(SCRIPTS / "build_deck.py"), "--spec", str(spec_dir / "deck.json"),
                        "--design", fx["design"], "--shapes", fx["shapes"], "--design-index", fx["design_index"],
                        "--shapes-index", fx["shapes_index"]], capture_output=True, encoding="utf-8", cwd=str(cwd))
    rep = json.loads(p.stdout)
    assert p.returncode == 0 and rep["warnings"] == [], rep
    slides = Presentation(rep["output"]).slides
    colours = []
    for i in range(3):
        pic = {s.name: s for s in slides[i].shapes}["Photo"]
        from io import BytesIO
        colours.append(Image.open(BytesIO(pic.image.blob)).convert("RGB").getpixel((5, 5)))
    assert colours == [(10, 120, 200), (200, 20, 20), (0, 200, 0)]
    assert not any(s.name == "Body" and s.has_text_frame and s.text_frame.text == "t" for s in slides[2].shapes)


def test_missing_relative_input_names_both_folders(fx, tmp_path):
    code, rep, _ = build(fx, tmp_path, {"slides": [
        {"slide": 3, "fill": {"Title": "x", "Body": "t", "Photo": {"path": "nowhere/x.png", "alt": "x"}}},
        {"slide": 3, "fill": {"Title": "y", "Body": "t", "Photo": {"path": "https://example.org/x.png", "alt": "x"}}},
    ]})
    assert code == 0
    missing, url = [w for w in rep["warnings"] if "could not fill 'Photo'" in w]
    assert "nowhere/x.png" in missing and str(tmp_path) in missing and "the folder of deck.json" in missing
    assert "the working directory" in missing and str(Path.cwd()) in missing
    assert "https://example.org/x.png" in url and "folder of deck.json" not in url  # URLs are not rewritten


def test_table_rows_grow_and_shrink(fx, tmp_path):
    code, rep, _ = build(fx, tmp_path, {"slides": [
        {"slide": 5, "fill": {"Title": "T", "Table": [["a", "b"], ["c", "d"], ["e", "f"], ["g", "h"], ["i", "j"]]}},
        {"slide": 5, "fill": {"Title": "T", "Table": [["a", "b"]]}},
    ]})
    assert code == 0
    prs = Presentation(rep["output"])
    t1 = {s.name: s for s in prs.slides[0].shapes}["Table"].table
    t2 = {s.name: s for s in prs.slides[1].shapes}["Table"].table
    assert len(t1.rows) == 5 and t1.cell(4, 1).text == "j"
    assert len(t2.rows) == 1


def test_component_insert_across_packages_with_chart_and_group(fx, tmp_path):
    code, rep, _ = build(fx, tmp_path, {"slides": [
        {"slide": 6, "fill": {"Title": "Roadmap"}, "insert": [
            {"source": "shapes", "slide": 1, "component": "Timeline 4 steps", "into": "Content Area",
             "fill": {"Step 1 title": "Discover", "Step 4 title": "Scale"}}]},
        {"slide": 6, "fill": {"Title": "Numbers", "Content Area": None}, "insert": [
            {"source": "shapes", "slide": 2, "component": "Bar Chart", "box": [5, 25, 90, 70],
             "fill": {"Bar Chart": {"categories": ["Q1", "Q2"], "series": {"2026": [4, 9]}}}}]},
        {"slide": 6, "fill": {"Title": "Second chart, untouched data", "Content Area": None}, "insert": [
            {"source": "shapes", "slide": 2, "component": "Bar Chart", "box": [5, 25, 90, 70]}]},
    ]})
    assert code == 0, rep
    assert not any("Content Area" in w for w in rep["warnings"])
    prs = Presentation(rep["output"])
    assert "Content Area" not in [s.name for s in prs.slides[0].shapes]
    grp = [s for s in prs.slides[0].shapes if s.name == "Timeline 4 steps"][0]
    texts = {s.name: s.text_frame.text for s in grp.shapes if s.has_text_frame}
    assert texts["Step 1 title"] == "Discover" and texts["Step 4 title"] == "Scale" and texts["Step 2 title"] == "Step 2"
    assert grp.width <= prs.slide_width  # scaled into the frame

    def chart_of(i):
        return [s for s in prs.slides[i].shapes if s.name == "Bar Chart"][0].chart
    assert list(chart_of(1).plots[0].categories) == ["Q1", "Q2"]
    assert list(chart_of(2).plots[0].categories) == ["A", "B", "C"]  # independent copy: not changed by slide 2
    names = zipfile.ZipFile(rep["output"]).namelist()
    assert len([n for n in names if n.startswith("ppt/charts/chart")]) == 2
    assert len([n for n in names if n.startswith("ppt/embeddings/")]) == 2


def test_warnings_for_unknown_key_overflow_and_missing_alt(fx, tmp_path):
    long = "word " * 400
    code, rep, _ = build(fx, tmp_path, {"slides": [{"slide": 2, "fill": {"Title": "T", "Body": long, "Nope": "x"}}]})
    assert code == 0
    assert any("no shape named 'Nope'" in w for w in rep["warnings"])
    assert any("overflow" in w for w in rep["warnings"])


def test_strict_mode_and_bad_slide_number(fx, tmp_path):
    code, rep, _ = build(fx, tmp_path, {"slides": [{"slide": 1, "fill": {"Typo": "x"}}]}, extra=["--strict"])
    assert code == 2 and "Typo" in rep["error"]
    code, rep, _ = build(fx, tmp_path, {"slides": [{"slide": 99}]})
    assert code == 2 and "does not exist" in rep["error"]


def test_null_deletes_shape(fx, tmp_path):
    code, rep, _ = build(fx, tmp_path, {"slides": [{"slide": 1, "fill": {"Title": "T", "Subtitle": None}}]})
    assert code == 0
    assert [s.name for s in Presentation(rep["output"]).slides[0].shapes] == ["Title"]


# ------------------------------------------------------------------ lint
def test_lint_flags_leftovers_and_passes_clean_deck(fx, tmp_path):
    code, rep, _ = build(fx, tmp_path, {"slides": [{"slide": 1, "fill": {"Title": "Real title", "Subtitle": "Real subtitle"}}]})
    c, out, _ = run("lint_deck.py", rep["output"], "--json")
    assert c == 0, out
    code, rep, _ = build(fx, tmp_path, {"slides": [{"slide": 1, "fill": {"Title": "Lorem ipsum dolor", "Subtitle": "{{todo}}"}}]})
    c, out, _ = run("lint_deck.py", rep["output"], "--json")
    assert c == 1
    assert len(json.loads(out)["errors"]) == 2


def test_inspect_and_emit_index(fx, tmp_path):
    c, out, _ = run("inspect_template.py", fx["design"], "--slide", "4")
    assert c == 0 and "'Number 1'" in out
    dest = tmp_path / "draft.xlsx"
    c, out, _ = run("inspect_template.py", fx["design"], "--emit-index", dest)
    assert c == 0
    entries = load_index(dest)
    assert len(entries) == 6 and entries[3]["field_list"][0]["name"] == "Title"


def test_read_index_cli(fx):
    c, out, _ = run("read_index.py", fx["design_index"], "--query", "table")
    assert c == 0 and "#5" in out
    c, out, _ = run("read_index.py", fx["design_index"], "--summary")
    assert "6 entries" in out


# ------------------------------------------------------------------ render
@pytest.mark.skipif(not (shutil.which("soffice") and shutil.which("pdftoppm")), reason="LibreOffice/poppler not installed")
def test_render_roundtrip(fx, tmp_path):
    code, rep, _ = build(fx, tmp_path, {"slides": [
        {"slide": 1, "fill": {"Title": "Render me", "Subtitle": "ok"}},
        {"slide": 6, "fill": {"Title": "Chart"}, "insert": [{"source": "shapes", "slide": 2, "component": "Bar Chart", "into": "Content Area"}]},
    ]})
    assert code == 0
    c, out, _ = run("render_check.py", rep["output"], "--out-dir", tmp_path / "r", "--dpi", 40, "--odp")
    assert c == 0, out
    info = json.loads(out)
    assert len(info["slides"]) == 2 and Path(info["contact_sheet"]).is_file()
    assert Path(info["odp"]).is_file() and Path(info["odp"]).suffix == ".odp"
