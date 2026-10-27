"""The logo slot: deck.json "logo" fills the "Logo" slot of the covers and the closing slide.

Logos are generated here (plain shapes in made-up colours), never taken from a real brand.
"""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / "skills" / "make-clean-pptx"
SCRIPTS = SKILL / "scripts"
ASSETS = SKILL / "assets"
sys.path.insert(0, str(SCRIPTS))

from _common import is_decorative  # noqa: E402
from read_index import load_index  # noqa: E402

EMU = 914400
SLOTS = {"L-01", "L-02", "L-03", "L-04", "L-05", "L-49"}
SLIDES = [
    {"id": "L-01", "fill": {"Title": "Annual review", "Subtitle": "Where we stand"}},
    {"id": "L-03", "fill": {"Title": "Annual review"}},
    {"id": "L-05", "fill": {"Title": "Annual review"}},
    {"id": "L-09", "fill": {"Title": "Chapter"}},
    {"id": "L-49", "fill": {"Contact 1 name": "Team"}},
]


@pytest.fixture(scope="module")
def logos(tmp_path_factory):
    d = tmp_path_factory.mktemp("brand")
    wide = Image.new("RGBA", (800, 200), (0, 0, 0, 0))
    draw = ImageDraw.Draw(wide)
    draw.rectangle([0, 0, 190, 199], fill=(30, 60, 140, 255))
    draw.ellipse([230, 20, 390, 180], fill=(230, 120, 40, 255))
    wide.save(d / "wide.png")
    Image.new("RGB", (300, 600), (30, 60, 140)).save(d / "tall.jpg")
    (d / "mark.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">'
                                '<circle cx="50" cy="50" r="45" fill="#1e3c8c"/>'
                                '<rect x="30" y="30" width="40" height="40" fill="#e67828"/></svg>', encoding="utf-8")
    (d / "broken.png").write_bytes(b"not an image at all")
    Image.new("RGB", (40, 40)).save(d / "anim.gif")
    return d


def build(tmp_path, extra=None, spec_dir=None, cwd=None, out="deck.pptx"):
    spec_dir = spec_dir or tmp_path
    spec_dir.mkdir(parents=True, exist_ok=True)
    spec = dict({"output": str(tmp_path / out), "slides": SLIDES}, **(extra or {}))
    path = spec_dir / "deck.json"
    path.write_text(json.dumps(spec), encoding="utf-8")
    p = subprocess.run([sys.executable, str(SCRIPTS / "build_deck.py"), "--spec", str(path),
                        "--layouts", str(ASSETS / "layouts.pptx"), "--layouts-index", str(ASSETS / "layouts.csv")],
                       capture_output=True, encoding="utf-8", cwd=str(cwd or tmp_path))
    return p.returncode, json.loads(p.stdout), tmp_path / out


def logo_shapes(deck):
    return {i: s for i, slide in enumerate(Presentation(deck).slides, 1) for s in slide.shapes if s.name == "Logo"}


def lint(deck):
    p = subprocess.run([sys.executable, str(SCRIPTS / "lint_deck.py"), str(deck), "--json"],
                       capture_output=True, encoding="utf-8")
    return json.loads(p.stdout)


def test_the_slots_are_in_the_index_with_their_role():
    got = {e["id"] for e in load_index(ASSETS / "layouts.csv") for f in e["field_list"] if f["name"] == "Logo"}
    assert got == SLOTS
    for e in load_index(ASSETS / "layouts.csv"):
        for f in e["field_list"]:
            if f["name"] == "Logo":
                assert f["optional"] and f["type"] == "image" and f.get("role") == "logo"


def test_no_logo_key_leaves_no_slot(tmp_path):
    code, rep, deck = build(tmp_path)
    assert code == 0 and rep["warnings"] == []
    assert logo_shapes(deck) == {}
    assert lint(deck)["warnings"] == []


def test_logo_fills_every_slot_and_keeps_its_aspect_ratio(tmp_path, logos):
    code, rep, deck = build(tmp_path, {"logo": str(logos / "wide.png"), "logo_alt": "Example Brand"})
    assert code == 0 and rep["warnings"] == [], rep
    shapes = logo_shapes(deck)
    assert sorted(shapes) == [1, 2, 3, 5]  # L-09 (slide 4) has no slot
    for s in shapes.values():
        assert abs(s.width / s.height - 4.0) < 0.01  # 800 x 200, not stretched to the 3.2:1 slot
        assert s.width <= 1.6 * EMU + 1 and s.height <= 0.5 * EMU + 1
        assert (s.crop_left, s.crop_right, s.crop_top, s.crop_bottom) == (0, 0, 0, 0)
        assert s._element.xpath(".//p:cNvPr")[0].get("descr") == "Example Brand"
    # Left corner on the margin, centred on the typographic cover, right corner on the right margin
    assert abs(shapes[1].left - 0.6 * EMU) < 2
    assert abs(shapes[2].left + shapes[2].width / 2 - 13.333 * EMU / 2) < EMU * 0.01
    assert abs(shapes[3].left + shapes[3].width - (13.333 - 0.6) * EMU) < EMU * 0.01
    assert lint(deck)["warnings"] == []


def test_tall_logo_without_alt_is_decorative(tmp_path, logos):
    code, rep, deck = build(tmp_path, {"logo": str(logos / "tall.jpg")})
    assert code == 0 and rep["warnings"] == [], rep
    for s in logo_shapes(deck).values():
        assert abs(s.width / s.height - 0.5) < 0.01 and abs(s.height - 0.5 * EMU) < 2
        assert is_decorative(s._element.xpath(".//p:cNvPr")[0])
    assert lint(deck)["warnings"] == []  # decorative: no "picture has no alt text"


def test_svg_logo_becomes_native_shapes_in_its_colours(tmp_path, logos):
    code, rep, deck = build(tmp_path, {"logo": str(logos / "mark.svg"), "logo_alt": "Example Brand"})
    assert code == 0 and rep["warnings"] == [], rep
    shapes = logo_shapes(deck)
    assert len(shapes) == 4
    for s in shapes.values():
        assert s.width == s.height
    from lxml import etree

    xml = etree.tostring(Presentation(deck).slides[0]._element).decode().upper()
    assert "1E3C8C" in xml and "<P:PIC" not in xml  # brand colours kept ("keep"), no picture


def test_relative_logo_path_is_read_next_to_the_spec_then_in_the_current_folder(tmp_path, logos):
    specs = tmp_path / "specs"
    specs.mkdir()
    (specs / "logo.png").write_bytes((logos / "wide.png").read_bytes())
    code, rep, deck = build(tmp_path, {"logo": "logo.png"}, spec_dir=specs, cwd=logos)
    assert code == 0 and len(logo_shapes(deck)) == 4
    code, rep, deck = build(tmp_path, {"logo": "tall.jpg"}, spec_dir=specs, cwd=logos, out="cwd.pptx")
    assert code == 0 and len(logo_shapes(deck)) == 4


def test_a_slide_can_drop_the_logo(tmp_path, logos):
    slides = [dict(SLIDES[0], fill=dict(SLIDES[0]["fill"], Logo=None))] + SLIDES[1:]
    code, rep, deck = build(tmp_path, {"logo": str(logos / "wide.png"), "slides": slides})
    assert code == 0 and sorted(logo_shapes(deck)) == [2, 3, 5]


def test_logo_without_slot_is_reported(tmp_path, logos):
    code, rep, _ = build(tmp_path, {"logo": str(logos / "wide.png"), "slides": [SLIDES[3]]})
    assert code == 0 and any("no slide has a logo slot" in w for w in rep["warnings"])


@pytest.mark.parametrize("name, words", [("missing.png", "not found"), ("broken.png", "not a readable image"),
                                         ("anim.gif", "PNG, JPG or SVG"), ("notes.txt", "PNG, JPG or SVG")])
def test_wrong_logo_file_is_a_plain_error(tmp_path, logos, name, words):
    code, rep, deck = build(tmp_path, {"logo": str(logos / name)})
    assert code == 2 and words in rep["error"] and "Traceback" not in rep["error"]
    assert not deck.exists()


def test_lint_warns_when_a_logo_covers_text(tmp_path, logos):
    code, rep, deck = build(tmp_path, {"logo": str(logos / "wide.png"), "logo_alt": "Example Brand"})
    prs = Presentation(deck)
    slide = prs.slides[0]
    title = next(s for s in slide.shapes if s.name == "Title")
    logo = next(s for s in slide.shapes if s.name == "Logo")
    logo.left, logo.top = title.left, title.top + title.height - logo.height
    prs.save(deck)
    found = [w for w in lint(deck)["warnings"] if w["shape"] == "Logo"]
    assert found and "overlaps text box 'Title'" in found[0]["message"]
