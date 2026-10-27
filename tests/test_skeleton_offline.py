"""Skeleton mode (no user content in the deck) and offline mode (no network at all)."""
from __future__ import annotations

import json
import re
import socket
import sys
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SKILL = ROOT / "skills" / "make-clean-pptx"
SCRIPTS = SKILL / "scripts"
ASSETS = SKILL / "assets"
sys.path.insert(0, str(SCRIPTS))

import build_deck  # noqa: E402
import fetch_image  # noqa: E402
import fetch_templates  # noqa: E402
import find_icon  # noqa: E402
import lint_deck  # noqa: E402
from _common import OFFLINE_ENV, iter_shapes  # noqa: E402
from read_index import load_index  # noqa: E402

SECRETS = ["Northwind merger", "salary grid", "48.7", "Q4 layoffs", "speaker secret", "Codename"]
SLIDES = [
    {"id": "L-01", "fill": {"Title": "Northwind merger", "Subtitle": "salary grid", "Date": None},
     "notes": "speaker secret"},
    {"id": "L-10", "fill": {"Number": "48.7", "Title": "Q4 layoffs", "Footer": "Confidential"}},
    {"id": "L-37", "fill": {"Title": "Codename", "KPI 1 value": "48.7", "KPI 1 detail": "salary grid"}},
    {"id": "L-20", "fill": {"Icon 1": {"icon": "lucide:rocket"}}},
    {"id": "L-42", "fill": {"Table": [["Codename", "x"]]}},
    {"id": "L-31", "fill": {"Title": "Q4 layoffs"},
     "insert": [{"source": "visuals", "id": "V-CMP-02", "component": "V-CMP-02", "into": "Content Area",
                 "fill": {"Chart": {"categories": ["Northwind"], "series": {"salary grid": [48.7]}}}}]},
]


def _build(tmp_path, slides=SLIDES, extra=(), **spec_extra):
    spec = dict({"output": str(tmp_path / "deck.pptx"), "properties": {"title": "Northwind merger", "author": "Codename"},
                 "slides": slides}, **spec_extra)
    path = tmp_path / "deck.json"
    path.write_text(json.dumps(spec), encoding="utf-8")
    return path


def _run_build(capsys, path, *extra):
    code = build_deck.main(["--spec", str(path), *extra])
    return code, json.loads(capsys.readouterr().out)


@pytest.fixture()
def no_network(monkeypatch, tmp_path):
    """Any attempt to open a connection fails the test; the cache is empty."""
    opened = []

    def refuse(*a, **k):
        opened.append(a)
        raise AssertionError("a network connection was attempted in offline mode")

    monkeypatch.setattr(socket, "socket", refuse)
    monkeypatch.setattr(socket, "create_connection", refuse)
    monkeypatch.setattr(socket, "getaddrinfo", refuse)
    monkeypatch.setenv("MAKE_CLEAN_PPTX_CACHE", str(tmp_path / "cache"))
    monkeypatch.setenv(OFFLINE_ENV, "0")  # restored after the test, whatever --offline sets
    return opened


@pytest.fixture()
def skeleton(tmp_path, capsys, no_network):
    code, rep = _run_build(capsys, _build(tmp_path), "--skeleton", "--offline")
    assert code == 0, rep
    return rep


def test_skeleton_deck_has_no_user_content(skeleton):
    blob = b"".join(zipfile.ZipFile(skeleton["output"]).read(n) for n in zipfile.ZipFile(skeleton["output"]).namelist())
    for secret in SECRETS:
        assert secret.encode() not in blob, secret
    prs = Presentation(skeleton["output"])
    assert prs.core_properties.title == "" and prs.core_properties.author == ""
    sheet = Path(skeleton["fill_sheet"]).read_text(encoding="utf-8")
    assert not any(s in sheet for s in SECRETS)
    assert skeleton["skeleton"] is True and skeleton["fields_to_fill"] > 10


def test_skeleton_placeholders_state_the_index_limits(skeleton):
    limits = {}
    for kind in ("layouts", "visuals"):
        for e in load_index(ASSETS / f"{kind}.csv"):
            for f in e["field_list"]:
                limits[(e["id"], f["name"])] = f["limit"]
    prs = Presentation(skeleton["output"])
    ids = [s["id"] for s in SLIDES]
    checked = 0
    for n, slide in enumerate(prs.slides):
        for shp in iter_shapes(slide):
            if not shp.has_text_frame:
                continue
            m = re.search(r"\bmax (\d+)\b", shp.text_frame.text)
            if m:
                assert int(m.group(1)) == limits[(ids[n], shp.name)], (n + 1, shp.name)
                assert shp.text_frame.text.startswith("[") and shp.text_frame.text.endswith("]")
                checked += 1
    assert checked >= 10
    # optional fields named in fill are kept, others removed; null deletes
    names = {s.name for s in iter_shapes(prs.slides[2])}
    assert "KPI 1 detail" in names and "KPI 2 detail" not in names
    assert "Date" not in {s.name for s in prs.slides[0].shapes}
    footer = next(s for s in prs.slides[1].shapes if s.name == "Footer")
    assert footer.text_frame.text == "Confidential"  # the label chosen up front is kept
    # chart: dummy series, same shape as the template
    chart = next(s for s in iter_shapes(prs.slides[5]) if s.has_chart).chart
    assert [s.name for s in chart.plots[0].series] == ["Dummy data — replace (1)", "Dummy data — replace (2)"]
    assert set(chart.plots[0].series[0].values) == {10}
    # speaker notes list each field and its limit; the user's notes are not kept
    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "Number: max 6 characters" in notes and "Title: max 53 characters" in notes
    rows = [ln for ln in Path(skeleton["fill_sheet"]).read_text(encoding="utf-8").splitlines() if ln.startswith("| 2 |")]
    assert "| 2 | Number | max 6 characters |" in rows


def test_lint_lists_what_is_left_only_when_present(skeleton, tmp_path, capsys):
    rep = lint_deck.lint(skeleton["output"])
    assert not rep["errors"], rep["errors"]
    todo = [w for w in rep["warnings"] if w["message"].startswith("still to fill")]
    assert len(todo) >= 10
    assert any("dummy data" in w["message"] for w in rep["warnings"])
    assert any(w["shape"] == "Table" for w in todo)

    (tmp_path / "real").mkdir()
    real = _build(tmp_path / "real", slides=[{"id": "L-10", "fill": {"Number": "12%", "Title": "Churn fell"}}])
    code, rep2 = _run_build(capsys, real)
    assert code == 0
    warns = lint_deck.lint(rep2["output"])["warnings"]
    assert not any("still to fill" in w["message"] or "dummy data" in w["message"] for w in warns)


def test_offline_build_never_opens_a_socket(tmp_path, capsys, no_network):
    slides = [{"id": "L-20", "fill": {"Title": "Pillars", "Icon 1": {"icon": "tabler:rocket"},
                                      "Icon 2": {"icon": "lucide:rocket"}}}]
    code, rep = _run_build(capsys, _build(tmp_path, slides=slides), "--offline")
    assert code == 0, rep
    assert any("offline mode" in w and "Icon 1" in w for w in rep["warnings"])
    assert not any("Icon 2" in w for w in rep["warnings"])  # bundled icon still works
    assert not no_network


def test_offline_icons_images_and_libraries(monkeypatch, capsys, no_network):
    assert find_icon.main(["rocket", "--source", "lucide,tabler,material,openmoji", "--offline"]) == 0
    out = json.loads(capsys.readouterr().out)
    assert out["results"] and all(r["source"] == "lucide" for r in out["results"])
    assert all("offline mode" in e for e in out["errors"].values()) and len(out["errors"]) == 3

    assert find_icon.main(["--get", "material:rocket", "--offline"]) == 3
    assert "offline mode" in json.loads(capsys.readouterr().out)["error"]

    assert fetch_image.main(["--query", "team meeting", "--offline"]) == 3
    assert "offline mode" in json.loads(capsys.readouterr().out)["error"]

    monkeypatch.setenv("MAKE_CLEAN_PPTX_REPO", "someone/elsewhere")  # would download without --offline
    assert fetch_templates.ensure_templates(remote=True)["source"] == "bundled"
    assert not no_network
